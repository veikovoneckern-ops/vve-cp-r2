"""
DOKUMENTE -- aus Markdown wird Word, PowerPoint oder eine Webseite.

Warum ein eigener Baustein: Auf dem Server gibt es weder pandoc noch
LibreOffice. Neo sollte am 01.10.2026 ein Word-Dokument liefern, hatte kein
Werkzeug dafuer -- und behauptete dann, es erstellt zu haben. Ein Weg, der
immer gleich funktioniert, ist besser als Improvisation im Shell.

Genutzt von Neo (Werkzeug dokument_erstellen) und vom Stab (Ergebnisse als
Word/PowerPoint herunterladen). python-docx und python-pptx stehen in
requirements.txt.
"""
from __future__ import annotations

import html as htmlmod
import re
import time
from io import BytesIO
from pathlib import Path

from .konfig import DATEN

EXPORT_DIR = DATEN / "exports"
EXPORT_DIR.mkdir(parents=True, exist_ok=True)
FORMATE = ("docx", "pptx", "html", "md")


def dateiname(titel: str, endung: str) -> str:
    s = re.sub(r"[^A-Za-z0-9ÄÖÜäöüß_-]+", "-", (titel or "dokument").strip())[:60].strip("-") or "dokument"
    return f"{s}-{time.strftime('%Y%m%d-%H%M')}.{endung}"


# ------------------------------------------------------------ Markdown lesen
def _bloecke(md: str) -> list[tuple[str, object]]:
    """Grobe, verlaessliche Zerlegung: Ueberschriften, Absaetze, Listen, Tabellen, Code."""
    zeilen = (md or "").replace("\r\n", "\n").split("\n")
    out: list[tuple[str, object]] = []
    i = 0
    while i < len(zeilen):
        z = zeilen[i]
        if z.strip().startswith("```"):
            code = []
            i += 1
            while i < len(zeilen) and not zeilen[i].strip().startswith("```"):
                code.append(zeilen[i])
                i += 1
            i += 1
            out.append(("code", "\n".join(code)))
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", z)
        if m:
            out.append((f"h{len(m.group(1))}", m.group(2).strip()))
            i += 1
            continue
        if re.match(r"^\s*\|.*\|\s*$", z) and i + 1 < len(zeilen) and re.match(r"^\s*\|[\s:|-]+\|\s*$", zeilen[i + 1]):
            kopf = [c.strip() for c in z.strip().strip("|").split("|")]
            i += 2
            reihen = []
            while i < len(zeilen) and re.match(r"^\s*\|.*\|\s*$", zeilen[i]):
                reihen.append([c.strip() for c in zeilen[i].strip().strip("|").split("|")])
                i += 1
            out.append(("tabelle", [kopf] + reihen))
            continue
        if re.match(r"^\s*[-*•]\s+", z):
            punkte = []
            while i < len(zeilen) and re.match(r"^\s*[-*•]\s+", zeilen[i]):
                punkte.append(re.sub(r"^\s*[-*•]\s+", "", zeilen[i]))
                i += 1
            out.append(("liste", punkte))
            continue
        if re.match(r"^\s*\d+[.)]\s+", z):
            punkte = []
            while i < len(zeilen) and re.match(r"^\s*\d+[.)]\s+", zeilen[i]):
                punkte.append(re.sub(r"^\s*\d+[.)]\s+", "", zeilen[i]))
                i += 1
            out.append(("nummern", punkte))
            continue
        if re.match(r"^\s*---+\s*$", z):
            out.append(("trenner", None))
            i += 1
            continue
        if not z.strip():
            i += 1
            continue
        absatz = []
        while i < len(zeilen) and zeilen[i].strip() and not re.match(r"^(#{1,4}\s|```|\s*[-*•]\s|\s*\d+[.)]\s|\s*\||\s*---+\s*$)", zeilen[i]):
            absatz.append(zeilen[i].strip())
            i += 1
        if absatz:
            out.append(("absatz", " ".join(absatz)))
        else:
            i += 1
    return out


_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\n]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _laeufe(text: str) -> list[tuple[str, str]]:
    """Text in (art, inhalt): normal, fett, kursiv, code, link."""
    teile = []
    for stueck in _INLINE.split(text):
        if not stueck:
            continue
        if stueck.startswith("**") and stueck.endswith("**"):
            teile.append(("fett", stueck[2:-2]))
        elif stueck.startswith("`") and stueck.endswith("`"):
            teile.append(("code", stueck[1:-1]))
        elif stueck.startswith("[") and "](" in stueck:
            t, u = stueck[1:-1].split("](", 1)
            teile.append(("link", f"{t} ({u})"))
        elif stueck.startswith("*") and stueck.endswith("*") and len(stueck) > 2:
            teile.append(("kursiv", stueck[1:-1]))
        else:
            teile.append(("normal", stueck))
    return teile


def _klartext(text: str) -> str:
    return "".join(t for _, t in _laeufe(text))


# ------------------------------------------------------------ Word
def zu_docx(titel: str, md: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    stil = doc.styles["Normal"]
    stil.font.name = "Calibri"
    stil.font.size = Pt(11)

    def absatz_mit(p, text):
        for art, t in _laeufe(text):
            r = p.add_run(t)
            if art == "fett":
                r.bold = True
            elif art == "kursiv":
                r.italic = True
            elif art == "code":
                r.font.name = "Consolas"
            elif art == "link":
                r.font.color.rgb = RGBColor(0x0F, 0x5E, 0x68)

    bl = _bloecke(md)
    if titel and not (bl and bl[0][0] == "h1"):
        doc.add_heading(titel, level=0)
    for art, inhalt in bl:
        if art.startswith("h"):
            doc.add_heading(_klartext(str(inhalt)), level=min(int(art[1]), 3) if art != "h1" else 1)
        elif art == "absatz":
            absatz_mit(doc.add_paragraph(), str(inhalt))
        elif art in ("liste", "nummern"):
            for punkt in inhalt:  # type: ignore[union-attr]
                absatz_mit(doc.add_paragraph(style="List Bullet" if art == "liste" else "List Number"), punkt)
        elif art == "tabelle":
            reihen = inhalt  # type: ignore[assignment]
            spalten = max(len(r) for r in reihen)
            t = doc.add_table(rows=0, cols=spalten)
            t.style = "Table Grid"
            for ri, reihe in enumerate(reihen):
                zellen = t.add_row().cells
                for ci in range(spalten):
                    zellen[ci].text = ""
                    p = zellen[ci].paragraphs[0]
                    absatz_mit(p, reihe[ci] if ci < len(reihe) else "")
                    if ri == 0:
                        for r in p.runs:
                            r.bold = True
            doc.add_paragraph()
        elif art == "code":
            p = doc.add_paragraph()
            r = p.add_run(str(inhalt))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif art == "trenner":
            doc.add_paragraph("")
    puffer = BytesIO()
    doc.save(puffer)
    return puffer.getvalue()


# ------------------------------------------------------------ PowerPoint
def folien(md: str) -> list[dict[str, object]]:
    """Folien: getrennt durch '---' oder durch Ueberschriften der Ebene 1/2."""
    teile = [t for t in re.split(r"(?m)^\s*---+\s*$", md or "") if t.strip()]
    if len(teile) <= 1:
        teile = [t for t in re.split(r"(?m)^(?=#{1,2}\s)", md or "") if t.strip()]
    out = []
    for t in teile:
        titel, punkte = "", []
        for art, inhalt in _bloecke(t):
            if art.startswith("h") and not titel:
                titel = _klartext(str(inhalt))
            elif art in ("liste", "nummern"):
                punkte += [_klartext(p) for p in inhalt]  # type: ignore[union-attr]
            elif art == "absatz":
                punkte.append(_klartext(str(inhalt)))
            elif art == "tabelle":
                punkte += [" · ".join(r) for r in inhalt]  # type: ignore[union-attr]
        out.append({"titel": titel or "Folie", "punkte": punkte[:8]})
    return out


def zu_pptx(titel: str, md: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    titelfolie = prs.slides.add_slide(prs.slide_layouts[0])
    titelfolie.shapes.title.text = titel or "Präsentation"
    if len(titelfolie.placeholders) > 1:
        titelfolie.placeholders[1].text = time.strftime("%d.%m.%Y")
    for f in folien(md):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = str(f["titel"])
        rahmen = s.placeholders[1].text_frame
        rahmen.clear()
        for i, punkt in enumerate(f["punkte"]):  # type: ignore[union-attr]
            p = rahmen.paragraphs[0] if i == 0 else rahmen.add_paragraph()
            p.text = str(punkt)
            p.font.size = Pt(20)
    puffer = BytesIO()
    prs.save(puffer)
    return puffer.getvalue()


# ------------------------------------------------------------ HTML
def zu_html(titel: str, md: str) -> str:
    teile = []
    for art, inhalt in _bloecke(md):
        def fmt(s: str) -> str:
            out = ""
            for a, t in _laeufe(s):
                t = htmlmod.escape(t)
                out += {"fett": f"<strong>{t}</strong>", "kursiv": f"<em>{t}</em>", "code": f"<code>{t}</code>"}.get(a, t)
            return out
        if art.startswith("h"):
            teile.append(f"<{art}>{fmt(str(inhalt))}</{art}>")
        elif art == "absatz":
            teile.append(f"<p>{fmt(str(inhalt))}</p>")
        elif art in ("liste", "nummern"):
            tag = "ul" if art == "liste" else "ol"
            teile.append(f"<{tag}>" + "".join(f"<li>{fmt(p)}</li>" for p in inhalt) + f"</{tag}>")  # type: ignore[union-attr]
        elif art == "tabelle":
            reihen = inhalt  # type: ignore[assignment]
            teile.append("<table><tr>" + "".join(f"<th>{fmt(c)}</th>" for c in reihen[0]) + "</tr>" +
                         "".join("<tr>" + "".join(f"<td>{fmt(c)}</td>" for c in r) + "</tr>" for r in reihen[1:]) + "</table>")
        elif art == "code":
            teile.append(f"<pre>{htmlmod.escape(str(inhalt))}</pre>")
    return ("<!doctype html><meta charset='utf-8'><meta name='viewport' content='width=device-width'>"
            f"<title>{htmlmod.escape(titel)}</title><style>body{{font:16px/1.6 system-ui;max-width:860px;margin:32px auto;padding:0 20px;color:#0B1826}}"
            "h1,h2,h3{color:#0F5E68;line-height:1.25}table{border-collapse:collapse;margin:12px 0}td,th{border:1px solid #d8dee6;padding:6px 10px;text-align:left}"
            "th{background:#eef3f4}pre{background:#f4f6f8;padding:12px;border-radius:6px;overflow:auto}</style>"
            f"<h1>{htmlmod.escape(titel)}</h1>" + "".join(teile))


def erstellen(titel: str, md: str, fmt: str) -> Path:
    fmt = fmt if fmt in FORMATE else "docx"
    ziel = EXPORT_DIR / dateiname(titel, fmt)
    if fmt == "docx":
        ziel.write_bytes(zu_docx(titel, md))
    elif fmt == "pptx":
        ziel.write_bytes(zu_pptx(titel, md))
    elif fmt == "html":
        ziel.write_text(zu_html(titel, md), encoding="utf-8")
    else:
        ziel.write_text(md, encoding="utf-8")
    return ziel
