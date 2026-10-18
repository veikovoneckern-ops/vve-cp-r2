"""
PRAESENTATIONS-STILE -- der "Skill" fuer das Aussehen von Folien (18.10.2026).

Veiko: "Wenn ich durch das Cockpit Praesentationen erstellen lasse, soll der Look
and Feel so sein [VVEORG PPT Master Vorlage]. Der Footer soll immer so sein. Und
wenn eine Praesentation erstellt werden soll, moechte ich waehlen koennen, ob es
in diesem Look and Feel erstellt wird oder andere Vorgaben gelten."

Ein Stil ist ein ORDNER unter cockpit/stile/<id>/ mit stil.json (Farben,
Schriften, Footer, eine Anweisung fuer Annie, ein Bildstil) und den Logos
(logo-hell.png fuer dunkle, logo-dunkel.png fuer helle Folien). Ein weiterer Stil
ist also ein weiterer Ordner, kein Code. Drei Wahlen gibt es:

  vveorg   die Master-Vorlage (Vorgabe, umstellbar in den Einstellungen)
  cockpit  das bisherige Design aus folien.py (Marineblau, Papierweiss, Bernstein)
  eigen    Veikos Vorgaben als Text; ein lokales Modell leitet daraus Farben,
           Schriften und Footer ab, das Layout der Vorlage bleibt

Gewaehlt wird je Ergebnis (Spalte ergebnisse.stil) und es laesst sich jederzeit
umschalten: das Aussehen wird bei jedem Abruf aus dem Markdown gerechnet, ein
anderer Stil braucht also keine neue Ausarbeitung.

PowerPoint und Vorschau entstehen aus DERSELBEN Elementliste (_folie): jede Folie
wird einmal als Flaechen, Kreise, Linien, Texte und Bilder in Punkt-Koordinaten
beschrieben (960 x 540, wie die Vorlage), und zwei kleine Ausgaben setzen sie um.
So koennen Vorschau und Datei nicht auseinanderlaufen.
"""
from __future__ import annotations

import base64
import copy
import html as htmlmod
import json
import re
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from typing import Any

STILE_DIR = Path(__file__).resolve().parent / "stile"
B, H = 960, 540                  # Punkt, wie die Vorlage (13,333 x 7,5 Zoll)
WAHLEN = ("vveorg", "cockpit", "eigen")
# Schriften, die auf jedem Windows mit Office da sind -- PowerPoint ersetzt fehlende still.
SICHERE_SCHRIFTEN = ("Segoe UI", "Segoe UI Semibold", "Calibri", "Cambria", "Arial", "Georgia", "Verdana",
                     "Tahoma", "Trebuchet MS", "Times New Roman", "Century Gothic", "Garamond", "Franklin Gothic Medium")


# ------------------------------------------------------------ Stile laden und waehlen
@lru_cache(maxsize=8)
def _ordner(sid: str) -> dict[str, Any]:
    p = STILE_DIR / sid / "stil.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.is_file() else {}


@lru_cache(maxsize=16)
def _logo(sid: str, art: str) -> bytes | None:
    p = STILE_DIR / sid / f"logo-{art}.png"
    return p.read_bytes() if p.is_file() else None


def liste() -> list[dict[str, str]]:
    """Fuer die Auswahl in der Oberflaeche."""
    v = _ordner("vveorg")
    return [{"id": "vveorg", "name": v.get("name") or "VVEORG", "beschreibung": v.get("beschreibung") or ""},
            {"id": "cockpit", "name": "Cockpit-Design", "beschreibung": "Das bisherige Design: Marineblau, Papierweiß, Bernstein."},
            {"id": "eigen", "name": "Eigene Vorgaben", "beschreibung": "Du beschreibst Farben, Schriften, Footer; das Layout der Vorlage bleibt."}]


def vorgabe() -> str:
    from . import db
    w = db.einstellung("praesentation_stil", "vveorg")
    return w if w in ("vveorg", "cockpit") else "vveorg"


def normal(stil: Any) -> dict[str, Any]:
    """Was am Ergebnis gespeichert wird: {"id": ..., ggf. "text", "werte"}. Leer = Vorgabe."""
    if isinstance(stil, str):
        stil = {"id": stil}
    if not isinstance(stil, dict) or stil.get("id") not in WAHLEN:
        return {"id": vorgabe()}
    return stil


def werte(stil: Any) -> dict[str, Any] | None:
    """Die Parameter fuer das Layout der Vorlage -- None heisst: das alte Cockpit-Design."""
    s = normal(stil)
    if s["id"] == "cockpit":
        return None
    w = copy.deepcopy(_ordner("vveorg"))
    if s["id"] == "eigen" and isinstance(s.get("werte"), dict):
        e = s["werte"]
        w["farben"].update({k: v for k, v in (e.get("farben") or {}).items() if k in w["farben"]})
        w["schriften"].update({k: v for k, v in (e.get("schriften") or {}).items() if k in w["schriften"]})
        w["fuss"].update({k: v for k, v in (e.get("fuss") or {}).items() if k in w["fuss"]})
        w["id"] = "eigen"
        w["bildstil"] = e.get("bildstil") or w.get("bildstil")
    return w


def anweisung(stil: Any) -> str:
    """Was Annie zusaetzlich zur Folienanweisung bekommt."""
    s = normal(stil)
    if s["id"] == "cockpit":
        return ""
    t = _ordner("vveorg").get("anweisung") or ""
    if s["id"] == "eigen" and s.get("text"):
        t += f"\nVEIKOS EIGENE VORGABEN FÜR DIESE PRÄSENTATION (Inhalt und Ton beachten; Farben, Schrift und Footer setzt das Cockpit): {s['text'][:800]}"
    return t


def bildstil(stil: Any) -> str:
    w = werte(stil)
    return (w or {}).get("bildstil") or ""


_HEX = re.compile(r"^[0-9A-Fa-f]{6}$")


async def aus_text(text: str) -> dict[str, Any]:
    """Veikos Beschreibung -> Farben, Schriften, Footer. Ein lokales Modell schreibt
    nur Werte in ein festes Schema; was nicht passt (kein Hex, unbekannte Schrift),
    faellt weg und die Vorlage gilt dort weiter -- lieber ein Teil der Vorgabe als
    ein kaputtes Deck."""
    from . import llm
    v = _ordner("vveorg")
    modell = await llm.modell_waehlen(llm.vorgabe("talk"), llm.vorgabe("arbeit"))
    if not modell:
        raise llm.ModellFehler("Kein lokales Modell installiert.")
    frage = ("Leite aus Veikos Vorgaben die Gestaltung einer Präsentation ab. Ändere nur, was die Vorgaben "
             "verlangen; alles andere lässt du weg (dann gilt die Vorlage).\n\nVORLAGE:\n"
             + json.dumps({k: v[k] for k in ("farben", "schriften", "fuss")}, ensure_ascii=False)
             + f"\n\nVORGABEN VON VEIKO:\n{text[:1500]}\n\n"
             "Farben als sechsstelliges Hex ohne #. 'dunkel' ist der Grund dunkler Folien, 'hell' der heller, "
             "'akzent' die Signalfarbe (Unterzeile, Linien), 'reihe' vier Farben für Nummern und Karten. "
             f"Schriften nur aus: {', '.join(SICHERE_SCHRIFTEN)}. "
             "fuss.logo true/false, fuss.name und fuss.web als Text (leer = weglassen). "
             "bildstil: ein Satz zur Bildsprache, falls die Vorgaben dazu etwas sagen.\n"
             'Antworte NUR mit JSON: {"farben": {...}, "schriften": {...}, "fuss": {...}, "bildstil": "..."}')
    d = llm.json_aus(await llm.chat(modell, "", [{"role": "user", "content": frage}], json_format=True,
                                    temperatur=0.2, denken=False, timeout=180)) or {}
    aus: dict[str, Any] = {"farben": {}, "schriften": {}, "fuss": {}}
    for k, x in (d.get("farben") or {}).items():
        if k in ("reihe", "reihe_hell") and isinstance(x, list):
            r = [str(c).lstrip("#") for c in x if _HEX.match(str(c).lstrip("#"))][:4]
            if len(r) == 4:
                aus["farben"][k] = r
        elif k in v["farben"] and _HEX.match(str(x).lstrip("#")):
            aus["farben"][k] = str(x).lstrip("#").upper()
    for k, x in (d.get("schriften") or {}).items():
        if k in v["schriften"] and str(x) in SICHERE_SCHRIFTEN:
            aus["schriften"][k] = str(x)
    f = d.get("fuss") or {}
    if isinstance(f.get("logo"), bool):
        aus["fuss"]["logo"] = f["logo"]
    for k in ("name", "web"):
        if isinstance(f.get(k), str):
            aus["fuss"][k] = f[k].strip()[:60]
    if isinstance(d.get("bildstil"), str) and d["bildstil"].strip():
        aus["bildstil"] = d["bildstil"].strip()[:300]
    return aus


# ------------------------------------------------------------ Bausteine einer Folie
def _flaeche(x, y, w, h, farbe, rund=0.0, rand=None, form="rect"):
    return {"art": "form", "form": form, "x": x, "y": y, "w": w, "h": h, "farbe": farbe, "rund": rund, "rand": rand}


def _linie(x, y, w, farbe, staerke=0.75):
    return {"art": "linie", "x": x, "y": y, "w": w, "farbe": farbe, "staerke": staerke}


def _lauf(t, groesse, farbe, schrift, fett=False, kursiv=False, abstand=0):
    return {"t": t, "g": groesse, "c": farbe, "s": schrift, "b": fett, "i": kursiv, "spc": abstand}


def _text(x, y, w, h, absaetze, anker="oben", ausr="links", zeile=1.1, aufz=None):
    """absaetze: Liste von Laeufen-Listen. aufz: Farbe eines Aufzaehlungszeichens."""
    return {"art": "text", "x": x, "y": y, "w": w, "h": h, "abs": absaetze, "anker": anker, "ausr": ausr,
            "zeile": zeile, "aufz": aufz}


def _bild(roh, x, y, w, h, rund=0.0):
    return {"art": "bild", "roh": roh, "x": x, "y": y, "w": w, "h": h, "rund": rund}


def _passend(n: int, groesse: float, klein: float, viel: int) -> float:
    """Je mehr Text, desto kleiner -- in Stufen, damit nichts aus der Folie laeuft."""
    return groesse if n <= viel else max(klein, round(groesse - (n - viel) * groesse / max(viel, 1) * 0.6, 1))


def _teilen(p: str) -> tuple[str, str]:
    """'Kurzer Titel — Erklaerung' -> Karte mit Ueberschrift; sonst nur Text."""
    m = re.match(r"^(.{3,60}?)\s+[—–-]\s+(.+)$", p) or re.match(r"^([^:]{3,50}):\s+(.+)$", p)
    return (m.group(1).strip(), m.group(2).strip()) if m else ("", p)


def _kopf(el, t, unter, dunkel, w, breite=880):
    F, S = w["farben"], w["schriften"]
    if t:
        el.append(_text(40, 14, breite, 50, [[_lauf(t, _passend(len(t), 28, 20, 55), F["titel_dunkel"] if dunkel else F["titel_hell"],
                                                      S["kopf"], fett=True)]], anker="unten", zeile=1.0))
    if unter:
        el.append(_text(41, 68, breite, 18, [[_lauf(unter.upper(), 10.5, F["akzent"] if dunkel else F["akzent_hell"], S["text"],
                                                    fett=True, abstand=300)]]))


def _fuss(el, w, dunkel, gross=False):
    F, S, fu = w["farben"], w["schriften"], w["fuss"]
    if not gross:
        el.append(_linie(40, 500, 881, F["linie_dunkel"] if dunkel else F["linie_hell"]))
    zeilen = []
    if fu.get("name"):
        zeilen.append([_lauf(fu["name"], 14 if gross else 10.5, F["titel_dunkel"] if dunkel else F["text_hell"], S["text"], fett=True)])
    if fu.get("web"):
        zeilen.append([_lauf(fu["web"], 11 if gross else 8.5, F["titel_dunkel"] if dunkel else F["grau_hell"], S["text"])])
    logo = _logo("vveorg", "hell" if dunkel else "dunkel") if fu.get("logo") else None
    lw, lh = (68, 43) if gross else (36, 23)
    rechts = 921 - (lw + 8 if logo else 0)
    if zeilen:
        el.append(_text(rechts - 220, 461 if gross else 503, 220, 50 if gross else 30, zeilen, ausr="rechts", zeile=1.0))
    if logo:
        el.append(_bild(logo, 921 - lw, 463 if gross else 505, lw, lh))


def _kreis(el, x, y, d, farbe, zahl, schrift, voll=False, zahlfarbe=None):
    el.append(_flaeche(x, y, d, d, farbe if voll else None, rand=None if voll else (farbe, 2), form="kreis"))
    el.append(_text(x, y, d, d, [[_lauf(zahl, d * (0.42 if voll else 0.38), zahlfarbe or ("FFFFFF" if voll else farbe), schrift, fett=True)]],
                    anker="mitte", ausr="mitte", zeile=1.0))


def _folie(f: dict[str, Any], deck_titel: str, roh: bytes | None, w: dict[str, Any]) -> tuple[bool, list[dict]]:
    """Eine Folie als Elementliste. Gibt (dunkel, elemente) zurueck."""
    F, S = w["farben"], w["schriften"]
    typ, t, unter = f["typ"], f["titel"] or deck_titel or "", f.get("unter") or ""
    dunkel = typ in ("titel", "kapitel", "aussage", "zitat", "abschluss") or (typ == "bild" and roh is not None)
    el: list[dict] = []
    td, tt, tg = (F["titel_dunkel"], F["text_dunkel"], F["grau_dunkel"]) if dunkel else (F["titel_hell"], F["text_hell"], F["grau_hell"])
    reihe = F["reihe"] if dunkel else F["reihe_hell"]

    if typ == "titel":
        breit = 470 if roh else 820
        if roh:
            el.append(_bild(roh, 540, 0, 420, 440))
        el.append(_text(48, 60, breit, 140, [[_lauf(t, _passend(len(t), 46, 30, 34), td, S["titelfolie"], fett=True)]],
                        anker="unten", zeile=1.0))
        el.append(_flaeche(52, 214, 86, 2.2, F["akzent"]))
        if unter:
            el.append(_text(52, 232, breit, 60, [[_lauf(unter, 18, tg, S["text"])]]))
        _fuss(el, w, True, gross=True)
        return True, el

    if typ == "bild" and roh is not None:
        el.append(_bild(roh, 0, 0, B, H))
        el.append(_flaeche(0, 0, B, 96, "000000", rund=0.0))
        el[-1]["alpha"] = 45
        el.append(_text(40, 22, 880, 40, [[_lauf(t.upper(), _passend(len(t), 22, 16, 40), "FFFFFF", S["kopf"], fett=True, abstand=250)]],
                        anker="mitte", zeile=1.0))
        if unter:
            el.append(_text(41, 62, 880, 18, [[_lauf(unter.upper(), 10.5, F["akzent"], S["text"], fett=True, abstand=300)]]))
        _fuss(el, w, True)
        return True, el

    if typ == "kapitel":
        el.append(_text(40, 120, 300, 120, [[_lauf(f"{f.get('nr', 1):02d}", 96, F["akzent"], S["akzent"], fett=True)]], zeile=1.0))
        el.append(_text(40, 250, 600 if roh else 860, 110, [[_lauf(t, _passend(len(t), 40, 28, 40), td, S["kopf"], fett=True)]], zeile=1.0))
        nach = unter or " · ".join((f["rest"] + f["punkte"])[:3])
        if nach:
            el.append(_text(41, 338, 600 if roh else 860, 40, [[_lauf(nach.upper(), 11, F["akzent"], S["text"], fett=True, abstand=300)]]))
        if roh:
            el.append(_bild(roh, 680, 104, 240, 360, rund=0.04))
        _fuss(el, w, True)
        return True, el

    if typ == "zitat":
        z = f["zitat"] or t
        el.append(_text(40, 70, 880, 330, [[_lauf("„" + z.strip("„“\"") + "“", _passend(len(z), 30, 20, 120), td, S["kopf"],
                                                  fett=True, kursiv=True)]], anker="mitte", zeile=1.2))
        if f["quelle"]:
            el.append(_text(42, 410, 860, 24, [[_lauf(f["quelle"], 11, F["akzent"], S["text"], fett=True, abstand=200)]]))
        _fuss(el, w, True)
        return True, el

    beleg = (f["rest"] + f["punkte"])[:2] if typ == "aussage" else []
    # Eine Aussage ohne Beleg steht nur EINMAL da: gross in der Karte, nicht zusaetzlich als Ueberschrift.
    _kopf(el, t if typ != "aussage" or beleg else "", unter, dunkel, w, 880)

    if typ == "aussage":
        kw = 520 if roh else 880
        el.append(_flaeche(40, 112, kw, 360, F["karte_dunkel"], rund=0.04))
        haupt = " ".join(beleg) if beleg else t
        # Die Ueberschrift steht oben; in der Karte die Aussage gross in Akzentfarbe (wie Folie 17 der Vorlage).
        el.append(_text(70, 132, kw - 60, 320, [[_lauf(haupt, _passend(len(haupt), 26, 17, 110), F["akzent"], S["text"], fett=True, kursiv=True)]],
                        anker="mitte", ausr="mitte", zeile=1.25))
        if roh:
            el.append(_bild(roh, 580, 112, 340, 360, rund=0.04))

    elif typ == "zahlen":
        zahlen = f["zahlen"][:4] or [("", p) for p in f["punkte"][:4]]
        n = max(1, len(zahlen))
        bw = (880 - 24 * (n - 1)) / n
        for k, (zahl, was) in enumerate(zahlen):
            x = 40 + k * (bw + 24)
            el.append(_flaeche(x, 116, bw, 250, F["karte_hell"], rund=0.05, rand=(F["rand_hell"], 0.75)))
            el.append(_text(x + 22, 134, bw - 44, 80, [[_lauf(zahl or "–", _passend(len(zahl), 46, 26, 7), reihe[k % 4], S["akzent"], fett=True)]],
                            anker="unten", zeile=1.0))
            el.append(_flaeche(x + 22, 224, 40, 2, reihe[k % 4]))
            el.append(_text(x + 22, 238, bw - 44, 118, [[_lauf(was, _passend(len(was), 15, 11.5, 80), tt, S["text"])]], zeile=1.2))
        rest = f["rest"][:1]
        if rest:
            el.append(_text(40, 390, 880, 50, [[_lauf(rest[0], 15, F["akzent_hell"], S["text"], fett=True, kursiv=True)]], zeile=1.15))

    elif typ in ("zwei_spalten", "vergleich"):
        farben = (F["titel_hell"], F["akzent_hell"]) if typ == "zwei_spalten" else ("3E8E5E", F["akzent_hell"])
        for k, sp in enumerate(f["spalten"][:2]):
            x = 40 + k * 452
            el.append(_flaeche(x, 116, 428, 356, F["karte_hell"], rund=0.04, rand=(F["rand_hell"], 0.75)))
            el.append(_text(x + 20, 132, 388, 50, [[_lauf(sp["titel"], 16, farben[k], S["text"], fett=True)]], anker="unten", zeile=1.05))
            ps = sp["punkte"][:5]
            gr = _passend(sum(len(p) for p in ps), 15, 11.5, 240)
            el.append(_text(x + 20, 194, 388, 266, [[_lauf(p, gr, tt, S["text"])] for p in ps], zeile=1.15, aufz=farben[k]))

    elif typ == "abschluss":
        ps = (f["punkte"] or f["rest"])[:4]
        fazit = f["rest"][0] if f["punkte"] and f["rest"] else ""
        n = max(1, len(ps))
        bw = (880 - 26 * (n - 1)) / n
        hoch = 290 if fazit else 350
        for k, p in enumerate(ps):
            x = 40 + k * (bw + 26)
            kt, kb = _teilen(p)
            el.append(_flaeche(x, 112, bw, hoch, F["karte_dunkel"], rund=0.04))
            _kreis(el, x + 26, 136, 52, reihe[k % 4], f"{k + 1:02d}", S["akzent"])
            el.append(_text(x + 26, 204, bw - 52, 56, [[_lauf(kt or kb, _passend(len(kt or kb), 16, 12, 40), td, S["akzent"], fett=True)]],
                            zeile=1.05))
            if kt:
                el.append(_text(x + 26, 266, bw - 52, hoch - 170, [[_lauf(kb, _passend(len(kb), 14, 11, 110), tt, S["text"])]], zeile=1.2))
        if fazit:
            el.append(_text(40, 418, 880, 60, [[_lauf(fazit, _passend(len(fazit), 17, 13, 120), F["akzent"], S["text"], fett=True, kursiv=True)]],
                            zeile=1.2))

    elif typ == "quellen":
        qs = f.get("quellen") or []
        haelfte = (len(qs) + 1) // 2
        for k, teil in enumerate((qs[:haelfte], qs[haelfte:])):
            if teil:
                el.append(_text(40 + k * 450, 110, 430, 370,
                                [[_lauf(f"{n + 1 + k * haelfte}. {q[0][:80]}  ", 12, tt, S["text"]),
                                  _lauf(re.sub(r"^https?://(www[.])?", "", q[1]).split("/")[0], 12, tg, S["text"])] for n, q in enumerate(teil)],
                                zeile=1.1))

    else:   # punkte (und bild ohne gerechnetes Bild)
        ps = (f["punkte"] or f["rest"])[:6]
        geteilt = [_teilen(p) for p in ps]
        if roh or len(ps) > 4 or not ps:
            # Nummerierte Zeilen (wie die rechte Haelfte von Folie 2 der Vorlage).
            breit = 470 if roh else 860
            n = max(1, len(ps))
            zh = min(70, 360 / n)
            for k, (kt, kb) in enumerate(geteilt):
                y = 116 + k * zh
                _kreis(el, 40, y, 26, F["kreis_hell"], str(k + 1), S["text"], voll=True)
                laeufe = [[_lauf(kt, 15, td, S["text"], fett=True)], [_lauf(kb, 12.5, tg, S["text"])]] if kt else \
                    [[_lauf(kb, _passend(len(kb), 17, 13, 90), tt, S["text"])]]
                el.append(_text(80, y - 2, breit - 40, zh - 4, laeufe, zeile=1.1))
            if roh:
                el.append(_bild(roh, 560, 112, 360, 360, rund=0.04))
        else:
            # Drei oder vier Karten nebeneinander (wie Folie 16 der Vorlage).
            n = len(ps)
            bw = (880 - 24 * (n - 1)) / n
            for k, (kt, kb) in enumerate(geteilt):
                x = 40 + k * (bw + 24)
                el.append(_flaeche(x, 116, bw, 270, F["karte_hell"], rund=0.04, rand=(F["rand_hell"], 0.75)))
                _kreis(el, x + 22, 138, 32, F["kreis_hell"], str(k + 1), S["text"], voll=True)
                y = 186
                if kt:
                    el.append(_text(x + 22, y, bw - 44, 48, [[_lauf(kt, _passend(len(kt), 17, 13, 34), td, S["text"], fett=True)]], zeile=1.05))
                    y = 240
                el.append(_text(x + 22, y, bw - 44, 376 - y, [[_lauf(kb, _passend(len(kb), 14.5 if kt else 17, 11.5, 100), tt, S["text"])]],
                                zeile=1.2))
    _fuss(el, w, dunkel)
    return dunkel, el


def folien_elemente(md: str, titel: str, eid: str | None, w: dict[str, Any]) -> list[dict[str, Any]]:
    from . import folien, folienbilder
    fl = folien.parse(md)
    deck = (fl[0]["titel"] if fl and fl[0]["typ"] == "titel" and fl[0]["titel"] else titel or "")[:80]
    aus = []
    for f in fl:
        roh = folienbilder.vorhanden(eid, f["bild"]) if eid and f["bild"] else None
        dunkel, el = _folie(f, deck, roh, w)
        aus.append({"typ": f["typ"], "dunkel": dunkel, "elemente": el, "notizen": f["notizen"], "bild_fehlt": bool(f["bild"] and not roh)})
    return aus


# ------------------------------------------------------------ PowerPoint
def zu_pptx(titel: str, md: str, eid: str | None, w: dict[str, Any]) -> bytes:
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt
    from PIL import Image

    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(B), Pt(H)
    leer = prs.slide_layouts[6]
    anker = {"oben": MSO_ANCHOR.TOP, "mitte": MSO_ANCHOR.MIDDLE, "unten": MSO_ANCHOR.BOTTOM}
    ausr = {"links": PP_ALIGN.LEFT, "mitte": PP_ALIGN.CENTER, "rechts": PP_ALIGN.RIGHT}

    for fo in folien_elemente(md, titel, eid, w):
        s = prs.slides.add_slide(leer)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(w["farben"]["dunkel"] if fo["dunkel"] else w["farben"]["hell"])
        for e in fo["elemente"]:
            if e["art"] == "form":
                form = MSO_SHAPE.OVAL if e["form"] == "kreis" else MSO_SHAPE.ROUNDED_RECTANGLE if e["rund"] else MSO_SHAPE.RECTANGLE
                r = s.shapes.add_shape(form, Pt(e["x"]), Pt(e["y"]), Pt(e["w"]), Pt(e["h"]))
                r.shadow.inherit = False
                if e["rund"]:
                    r.adjustments[0] = e["rund"]
                if e["farbe"]:
                    r.fill.solid()
                    r.fill.fore_color.rgb = RGBColor.from_string(e["farbe"])
                    if e.get("alpha") is not None:
                        etree.SubElement(r.fill._xPr.find(f".//{A}srgbClr"), A + "alpha").set("val", str(e["alpha"] * 1000))
                else:
                    r.fill.background()
                if e["rand"]:
                    r.line.color.rgb = RGBColor.from_string(e["rand"][0])
                    r.line.width = Pt(e["rand"][1])
                else:
                    r.line.fill.background()
                r.text_frame.text = ""
            elif e["art"] == "linie":
                r = s.shapes.add_connector(1, Pt(e["x"]), Pt(e["y"]), Pt(e["x"] + e["w"]), Pt(e["y"]))
                r.line.color.rgb = RGBColor.from_string(e["farbe"])
                r.line.width = Pt(e["staerke"])
            elif e["art"] == "bild":
                b = Image.open(BytesIO(e["roh"]))
                p = s.shapes.add_picture(BytesIO(e["roh"]), Pt(e["x"]), Pt(e["y"]), Pt(e["w"]), Pt(e["h"]))
                q, z = b.width / b.height, e["w"] / e["h"]
                if abs(q - z) > 0.01:   # flaechenfuellend zuschneiden
                    if q > z:
                        p.crop_left = p.crop_right = (1 - z / q) / 2
                    else:
                        p.crop_top = p.crop_bottom = (1 - q / z) / 2
                if e["rund"]:
                    g = p._element.spPr.find(A + "prstGeom")
                    g.set("prst", "roundRect")
                    av = g.find(A + "avLst")
                    etree.SubElement(av, A + "gd", name="adj", fmla=f"val {int(e['rund'] * 100000)}")
            elif e["art"] == "text":
                tb = s.shapes.add_textbox(Pt(e["x"]), Pt(e["y"]), Pt(e["w"]), Pt(e["h"]))
                tf = tb.text_frame
                tf.word_wrap = True
                tf.vertical_anchor = anker[e["anker"]]
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                for i, laeufe in enumerate(e["abs"]):
                    par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    par.alignment = ausr[e["ausr"]]
                    par.line_spacing = e["zeile"]
                    if e["aufz"]:
                        g = laeufe[0]["g"] if laeufe else 14
                        ppr = par._p.get_or_add_pPr()
                        einzug = int(Pt(g * 1.1))
                        ppr.set("marL", str(einzug))
                        ppr.set("indent", str(-einzug))
                        clr = etree.SubElement(ppr, A + "buClr")
                        etree.SubElement(clr, A + "srgbClr").set("val", e["aufz"])
                        etree.SubElement(ppr, A + "buSzPct").set("val", "100000")
                        etree.SubElement(ppr, A + "buFont").set("typeface", "Arial")
                        etree.SubElement(ppr, A + "buChar").set("char", "•")
                        par.space_after = Pt(g * 0.5)
                    elif len(e["abs"]) > 1:
                        par.space_after = Pt(4)
                    for l in laeufe:
                        r = par.add_run()
                        r.text = l["t"]
                        r.font.size, r.font.name, r.font.bold, r.font.italic = Pt(l["g"]), l["s"], l["b"], l["i"]
                        r.font.color.rgb = RGBColor.from_string(l["c"])
                        if l["spc"]:
                            r.font._rPr.set("spc", str(l["spc"]))
        if fo["notizen"]:
            s.notes_slide.notes_text_frame.text = "\n".join(fo["notizen"])
    puffer = BytesIO()
    prs.save(puffer)
    return puffer.getvalue()


# ------------------------------------------------------------ Vorschau (HTML)
def _uri(roh: bytes) -> str:
    art = "png" if roh[:4] == b"\x89PNG" else "jpeg"
    return f"data:image/{art};base64," + base64.b64encode(roh).decode()


def zu_html(titel: str, md: str, eid: str | None, w: dict[str, Any]) -> str:
    """Dieselben Elemente wie in der PowerPoint, absolut gesetzt auf 960 x 540 Punkt."""
    teile = []
    for fo in folien_elemente(md, titel, eid, w):
        innen = []
        for e in fo["elemente"]:
            pos = f"left:{e['x']}px;top:{e['y']}px;width:{e['w']}px;"
            if e["art"] == "form":
                st = pos + f"height:{e['h']}px;"
                if e["farbe"]:
                    a = e["alpha"] / 100 if e.get("alpha") is not None else 1   # wie <a:alpha>: Deckkraft in Prozent
                    r, g, b = (int(e["farbe"][i:i + 2], 16) for i in (0, 2, 4))
                    st += f"background:rgba({r},{g},{b},{a});"
                if e["rand"]:
                    st += f"border:{e['rand'][1]}px solid #{e['rand'][0]};box-sizing:border-box;"
                if e["form"] == "kreis":
                    st += "border-radius:50%;"
                elif e["rund"]:
                    st += f"border-radius:{round(min(e['w'], e['h']) * e['rund'])}px;"
                innen.append(f"<div class='e' style='{st}'></div>")
            elif e["art"] == "linie":
                innen.append(f"<div class='e' style='{pos}height:0;border-top:{e['staerke']}px solid #{e['farbe']}'></div>")
            elif e["art"] == "bild":
                st = pos + f"height:{e['h']}px;object-fit:cover;"
                if e["rund"]:
                    st += f"border-radius:{round(min(e['w'], e['h']) * e['rund'])}px;"
                innen.append(f"<img class='e' alt='' style='{st}' src='{_uri(e['roh'])}'>")
            elif e["art"] == "text":
                jc = {"oben": "flex-start", "mitte": "center", "unten": "flex-end"}[e["anker"]]
                ta = {"links": "left", "mitte": "center", "rechts": "right"}[e["ausr"]]
                ps = []
                for laeufe in e["abs"]:
                    spans = "".join(
                        f"<span style=\"font:{'italic ' if l['i'] else ''}{700 if l['b'] else 400} {l['g']}px/{e['zeile']} '{l['s']}',"
                        f"'Segoe UI',Calibri,sans-serif;color:#{l['c']};letter-spacing:{l['spc'] / 100}px\">{htmlmod.escape(l['t'])}</span>"
                        for l in laeufe)
                    if e["aufz"]:
                        ps.append(f"<p class='az' style='--bc:#{e['aufz']}'>{spans}</p>")
                    else:
                        ps.append(f"<p>{spans}</p>")
                innen.append(f"<div class='e t' style='{pos}height:{e['h']}px;justify-content:{jc};text-align:{ta}'>{''.join(ps)}</div>")
        grund = w["farben"]["dunkel"] if fo["dunkel"] else w["farben"]["hell"]
        hinweis = "<div class='kommt'>Bild wird gerechnet …</div>" if fo["bild_fehlt"] else ""
        notiz = (f"<details class='notiz'><summary>Sprechernotizen</summary>{''.join(f'<p>{htmlmod.escape(n)}</p>' for n in fo['notizen'])}</details>"
                 if fo["notizen"] else "")
        teile.append(f"<section style='background:#{grund}'>{''.join(innen)}{hinweis}</section>{notiz}")
    css = ("html,body{margin:0;background:#D9DCE1;overflow-x:hidden}.buehne{width:960px;transform-origin:0 0;padding:12px 0}"
           "section{width:960px;height:540px;position:relative;overflow:hidden;margin:0 0 8px;border-radius:8px;"
           "box-shadow:0 4px 18px rgba(12,20,32,.18)}.e{position:absolute}.t{display:flex;flex-direction:column;overflow:hidden}"
           ".t p{margin:0 0 4px}.t p:last-child{margin-bottom:0}.az{position:relative;padding-left:1.1em;margin-bottom:.5em!important}"
           ".az:before{content:'•';position:absolute;left:0;color:var(--bc);font:700 1em Arial}"
           ".az span{font-size:inherit}.kommt{position:absolute;right:10px;top:10px;font:12px system-ui;color:#8A94A6;"
           "background:rgba(255,255,255,.8);padding:3px 8px;border-radius:6px}"
           ".notiz{width:960px;margin:0 0 18px;font:13px/1.5 system-ui,sans-serif;color:#5B6878}.notiz summary{cursor:pointer}.notiz p{margin:4px 0}")
    return ("<!doctype html><meta charset='utf-8'><meta name='viewport' content='width=device-width'>"
            f"<title>{htmlmod.escape(titel)}</title><style>{css}</style><div class='buehne'>" + "".join(teile) + "</div><script>"
            "function passen(){var b=document.querySelector('.buehne');var w=document.documentElement.clientWidth;"
            "var s=Math.min(1.8,(w-24)/960);b.style.transform='scale('+s+')';b.style.marginLeft=((w-960*s)/2)+'px';"
            "document.body.style.height=(b.offsetHeight*s)+'px';}passen();window.addEventListener('resize',passen);</script>")
