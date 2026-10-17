"""
PRAESENTATIONEN MIT ECHTEM DESIGN (17.10.2026).

Veiko: "PPT und Bildgenerierung -- da bin ich wirklich nicht zufrieden, das soll
so gut vom Ergebnis sein, als ob ich es mit Claude machen wuerde." Befund an der
ExO-Praesentation vom 16.10.: Annie beschrieb eine Gestaltung ("dunkler
Hintergrund, grosses Serifen-Zitat zentriert, viel Schwarzraum") -- und die
PowerPoint wurde trotzdem eine weisse Standardfolie mit Calibri-Stichpunkten.
Ein Design gab es nicht; die Gestaltungsabsicht wurde weggeworfen.

Jetzt: Annie waehlt je Folie einen FOLIENTYP (Zeile "Typ: …"), und dieses Modul
setzt jeden Typ in einem festen, durchgestalteten Design um -- fuer PowerPoint
UND Vorschau aus DERSELBEN Zerlegung (parse), damit beide nie auseinanderlaufen.
Der Inhalt bleibt Markdown: Team-Chat, Fassungen, Daniels Pruefung und die
Folienbilder (folienbilder.py) arbeiten unveraendert weiter.

Folientypen und ihre Schreibweise (das steht so auch in ANWEISUNG):
  titel        # Titel / Untertitel: … / Bild: …
  kapitel      # Kapitelname (Nummer zaehlt das Modul selbst)
  aussage      # die Kernaussage als Satz / ein Satz Beleg darunter
  punkte       # Aussage-Ueberschrift / - drei bis fuenf kurze Punkte   (Vorgabe)
  zahlen       # Ueberschrift / - 42 % | was die Zahl bedeutet   (2 bis 4)
  zwei_spalten # Ueberschrift / ## Links-Titel / - … / ## Rechts-Titel / - …
  vergleich    wie zwei_spalten, links "vorher", rechts "nachher"
  zitat        > das Zitat / — Quelle
  bild         # Ueberschrift ueber einem Vollbild / Bild: …
  abschluss    # Schlussgedanke / - naechste Schritte
Dazu je Folie optional "Bild: …" (Beschreibung; folienbilder.py macht daraus ein
Bild) und "Notizen: …" (Sprechernotizen, in PowerPoint unter der Folie).
Alte Folien ohne "Typ:" werden erraten (erste Folie = titel, ">" = zitat,
zwei "##" = zwei_spalten, sonst punkte).
"""
from __future__ import annotations

import html as htmlmod
import re
import time
from io import BytesIO
from typing import Any

TYPEN = ("titel", "kapitel", "aussage", "punkte", "zahlen", "zwei_spalten", "vergleich", "zitat", "bild", "abschluss", "quellen")
# Quellen der Internetsuche (stab.netz_fuss) und Link-Zeilen sind kein Folieninhalt:
# sie standen am 17.10. als rohe Links in den Stichpunkten der letzten Folie. Sie
# werden gesammelt und bekommen am Ende eine eigene, kleine Quellenfolie.
_QUELLEN_KOPF = re.compile(r"^[_*\s]*(quellen|sources|literatur)[^:]{0,30}:?[_*\s]*$", re.I)
_LINK = re.compile(r"^\s*(?:[-*•]|\d+[.)])?\s*\[([^\]]+)\]\((https?://[^)\s]+)\)\s*$")

# Das Design: ruhiges Marineblau, warmes Papierweiss, Bernstein als Akzent,
# Petrol als zweite Farbe -- dieselben Toene wie das Cockpit. Schriften, die auf
# jedem Windows-Rechner da sind (PowerPoint ersetzt fehlende Schriften still).
DUNKEL, PAPIER, AKZENT, AKZENT2, TEXT, GRAU, LINIE = "0B1826", "F7F5F0", "C9861A", "0F5E68", "1E2A36", "5B6878", "D9D4C7"
KOPF_SCHRIFT, TEXT_SCHRIFT, ZITAT_SCHRIFT = "Segoe UI Semibold", "Segoe UI", "Georgia"

ANWEISUNG = (
    "Liefere die Präsentation als Folien in Markdown; trenne Folien mit einer Zeile, die nur '---' enthält. "
    "Jede Folie beginnt mit einer Zeile 'Typ: <folientyp>'. Folientypen und Schreibweise:\n"
    "- titel: '# Titel', 'Untertitel: …', dazu eine Bild-Zeile.\n"
    "- kapitel: '# Kapitelname' (trennt Abschnitte; die Nummer setzt das Cockpit).\n"
    "- aussage: '# die Kernaussage als ganzer Satz', darunter EIN Satz als Beleg.\n"
    "- punkte: '# Überschrift als Aussage', darunter 3 bis 5 Punkte '- …' mit höchstens 12 Wörtern.\n"
    "- zahlen: '# Überschrift', darunter 2 bis 4 Zeilen '- 42 % | was die Zahl bedeutet' (nur belegte Zahlen).\n"
    "- zwei_spalten: '# Überschrift', dann '## Titel links' mit Punkten und '## Titel rechts' mit Punkten.\n"
    "- vergleich: wie zwei_spalten, links das Heute, rechts das Morgen.\n"
    "- zitat: '> das Zitat' und '— Name, Quelle' -- NUR ein Zitat, das wörtlich im Material oder in den Quellen steht. "
    "Ein bekannter Satz ohne gesicherten Urheber ist kein Zitat (17.10.: ein Sinnspruch stand als Sam Altman da).\n"
    "- bild: '# kurze Überschrift' über einem Vollbild, dazu eine Bild-Zeile.\n"
    "- abschluss: '# Schlussgedanke', darunter 2 bis 4 nächste Schritte.\n"
    "REGELN WIE IN EINER GUTEN BERATER-PRÄSENTATION: Die Überschrift ist die Aussage der Folie als ganzer Satz, nicht "
    "ein Thema („Agenten verdreifachen den Durchsatz“, nicht „Agenten“). Höchstens 30 Wörter Text je Folie -- was du "
    "dazu sagen würdest, gehört in eine Zeile 'Notizen: …' (Sprechernotizen). Wechsle die Folientypen ab; nie drei "
    "punkte-Folien hintereinander. Erste Folie titel, letzte abschluss, bei mehr als acht Folien kapitel-Folien.\n"
    "BILDER: Titelfolie, bild-Folien und etwa jede dritte weitere Folie bekommen am Ende GENAU EINE Zeile "
    "'Bild: <Motiv, Stil, Licht, Farben, Stimmung>'. Beschreibe ein Motiv, keine Schrift, keine Logos, keine echten "
    "Personen; alle Bilder einer Präsentation im selben Stil. Gestaltungsanweisungen (Layout, Schrift, Farben der "
    "Folie) schreibst du NICHT -- das Design setzt das Cockpit.")

_TYP = re.compile(r"^\s*(?:typ|folientyp|layout)\s*:\s*([a-zäöü_ -]+)\s*$", re.I)
_UNTER = re.compile(r"^\s*untertitel\s*:\s*(.+)$", re.I)
_NOTIZ = re.compile(r"^\s*(?:notizen?|sprechernotiz(?:en)?|sprechertext)\s*:\s*(.*)$", re.I)
_ZAHL = re.compile(r"^\s*\**\s*([~≈+\-−]?\s*[\d.,]+\s*(?:%|x|×|mio\.?|mrd\.?|tsd\.?|k|m|bn|€|\$|jahre?|monate?|tage?|std\.?)?)\s*\**\s*(?:\||–|—|:|-)\s*(.+)$", re.I)


def _rein(s: str) -> str:
    return re.sub(r"\*\*|__", "", s).strip()


def parse(md: str) -> list[dict[str, Any]]:
    from . import folienbilder
    from .stab import HINWEIS_RE, _ohne_md
    folien: list[dict[str, Any]] = []
    quellen: list[tuple[str, str]] = []
    for nr, roh in enumerate(folienbilder.folien_roh(md)):
        f: dict[str, Any] = {"typ": "", "titel": "", "unter": "", "punkte": [], "spalten": [], "zahlen": [], "zitat": "",
                             "quelle": "", "notizen": [], "rest": [], "bild": folienbilder.bild_prompt(roh)}
        spalte = None
        notiz_folgt = False
        quellen_folgen = False
        for z in roh.splitlines():
            s = z.strip()
            if not s:
                notiz_folgt = False
                continue
            if _QUELLEN_KOPF.match(s):
                quellen_folgen = True
                continue
            if m := _LINK.match(s):
                quellen.append((_rein(m.group(1)), m.group(2)))
                continue
            if quellen_folgen and re.search(r"https?://", s):
                u = re.search(r"https?://[^\s)\]]+", s).group(0)
                quellen.append((_rein(re.sub(r"\[|\]|\(https?://[^)]+\)|https?://\S+", "", s)).strip(" -–:") or u, u))
                continue
            if m := _TYP.match(s):
                t = m.group(1).strip().lower().replace(" ", "_").replace("-", "_")
                f["typ"] = {"zwei_spalten": "zwei_spalten", "zweispaltig": "zwei_spalten", "zahl": "zahlen", "kennzahlen": "zahlen",
                            "titelfolie": "titel", "abschnitt": "kapitel", "kernaussage": "aussage", "statement": "aussage",
                            "quote": "zitat", "vollbild": "bild", "schluss": "abschluss", "fazit": "abschluss"}.get(t, t)
                continue
            if m := _NOTIZ.match(_ohne_md(s)):
                f["notizen"].append(m.group(1).strip())
                notiz_folgt = True
                continue
            if notiz_folgt and not s.startswith(("#", "-", "*", ">")):
                f["notizen"].append(s)
                continue
            if m := _UNTER.match(_ohne_md(s)):
                f["unter"] = _rein(m.group(1))
                continue
            if folienbilder.BILD_ZEILE.match(z) or HINWEIS_RE.match(_ohne_md(s.lstrip("-*• "))):
                continue                               # Regie, kein Folieninhalt
            if s.startswith("## "):
                spalte = {"titel": _rein(s[3:]), "punkte": []}
                f["spalten"].append(spalte)
                continue
            if s.startswith("#"):
                f["titel"] = f["titel"] or _rein(s.lstrip("#"))
                continue
            if s.startswith(">"):
                f["zitat"] = (f["zitat"] + " " + _rein(s.lstrip("> "))).strip().strip("„“\"")
                continue
            if s.startswith(("—", "–", "--")) and f["zitat"]:
                f["quelle"] = _rein(s.lstrip("—–- "))
                continue
            if s.startswith(("-", "*", "•")) or re.match(r"^\d+[.)]\s", s):
                p = _rein(re.sub(r"^(?:[-*•]|\d+[.)])\s*", "", s))
                if spalte is not None:
                    spalte["punkte"].append(p)
                else:
                    f["punkte"].append(p)
                continue
            f["rest"].append(_rein(s))
        # Zahlen erkennen (auch ohne Typ, wenn jede Zeile so aussieht)
        if f["typ"] == "zahlen" or (not f["typ"] and f["punkte"] and all(_ZAHL.match(p) for p in f["punkte"])):
            for p in f["punkte"]:
                m = _ZAHL.match(p)
                f["zahlen"].append((m.group(1).strip(), m.group(2).strip()) if m else ("", p))
        if not f["typ"]:
            f["typ"] = ("titel" if nr == 0 else "zitat" if f["zitat"] else "zwei_spalten" if len(f["spalten"]) >= 2
                        else "zahlen" if f["zahlen"] else "aussage" if not f["punkte"] and len(f["rest"]) <= 1 and f["titel"] else "punkte")
        if f["typ"] not in TYPEN:
            f["typ"] = "punkte"
        if f["typ"] in ("zwei_spalten", "vergleich") and len(f["spalten"]) < 2:
            f["typ"] = "punkte"
            f["punkte"] += [p for sp in f["spalten"] for p in sp["punkte"]]
        if f["typ"] == "zitat" and not f["zitat"]:
            f["zitat"] = " ".join(f["rest"] or f["punkte"][:1])
        if f["typ"] == "titel" and not f["unter"] and f["rest"]:
            f["unter"] = f["rest"].pop(0)
        folien.append(f)
    if quellen:
        folien.append({"typ": "quellen", "titel": "Quellen", "unter": "", "punkte": [], "spalten": [], "zahlen": [], "zitat": "",
                       "quelle": "", "notizen": [], "rest": [], "bild": None, "quellen": list(dict.fromkeys(quellen))[:24]})
    k = 0
    for f in folien:
        if f["typ"] == "kapitel":
            k += 1
            f["nr"] = k
    return folien


def _schrift(n: int, groesse: float, klein: float, viel: int) -> float:
    """Je mehr Text, desto kleiner -- in Stufen, damit nichts aus der Folie laeuft."""
    return groesse if n <= viel else max(klein, groesse - (n - viel) * 1.5)


# ------------------------------------------------------------ PowerPoint
def zu_pptx(titel: str, md: str, eid: str | None = None) -> bytes:
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    from . import folienbilder

    folien = parse(md)
    # In der Fusszeile der Titel der Praesentation, nicht der (oft lange) Titel des Ergebnisses.
    deck = (folien[0]["titel"] if folien and folien[0]["typ"] == "titel" and folien[0]["titel"] else titel or "")[:80]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    B, H = prs.slide_width, prs.slide_height
    leer = prs.slide_layouts[6]

    def farbe(hexwert: str) -> RGBColor:
        return RGBColor.from_string(hexwert)

    def flaeche(s, x, y, w, h, hexwert: str, alpha: int | None = None, form=MSO_SHAPE.RECTANGLE):
        r = s.shapes.add_shape(form, x, y, w, h)
        r.fill.solid()
        r.fill.fore_color.rgb = farbe(hexwert)
        r.line.fill.background()
        r.shadow.inherit = False
        if alpha is not None:   # Deckkraft: python-pptx kann es nicht direkt, also im XML
            sf = r.fill._xPr.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
            a = etree.SubElement(sf, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha")
            a.set("val", str(alpha * 1000))
        return r

    def aufzaehlung(p, groesse: float, zeichenfarbe: str):
        """Echtes Aufzaehlungszeichen mit haengendem Einzug -- vorher stand "■" als
        Text vorn, und jede zweite Zeile eines Punkts rutschte unter das Zeichen."""
        A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        ppr = p._p.get_or_add_pPr()
        einzug = int(Pt(groesse * 1.25))
        ppr.set("marL", str(einzug))
        ppr.set("indent", str(-einzug))
        for alt in ppr.findall(A + "buChar") + ppr.findall(A + "buClr") + ppr.findall(A + "buSzPct") + ppr.findall(A + "buFont"):
            ppr.remove(alt)
        clr = etree.SubElement(ppr, A + "buClr")
        etree.SubElement(clr, A + "srgbClr").set("val", zeichenfarbe)
        etree.SubElement(ppr, A + "buSzPct").set("val", "75000")
        etree.SubElement(ppr, A + "buFont").set("typeface", "Arial")
        etree.SubElement(ppr, A + "buChar").set("char", "■")
        p.space_after = Pt(groesse * 0.55)

    def text(s, x, y, w, h, inhalt, groesse: float, hexwert: str = TEXT, schrift: str = TEXT_SCHRIFT, fett=False,
             kursiv=False, ausrichtung=PP_ALIGN.LEFT, anker=MSO_ANCHOR.TOP, abstand: float = 1.1, punkte: str | None = None):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anker
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        zeilen = inhalt if isinstance(inhalt, list) else [inhalt]
        for i, z in enumerate(zeilen):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ausrichtung
            p.line_spacing = abstand
            if punkte:
                aufzaehlung(p, groesse, punkte)
            teile = z if isinstance(z, list) else [(z, hexwert, fett)]
            for t, c, f in teile:
                r = p.add_run()
                r.text = t
                r.font.size, r.font.name, r.font.bold, r.font.italic = Pt(groesse), schrift, f, kursiv
                r.font.color.rgb = farbe(c)
        return tb

    def bild(s, roh: bytes, x, y, w, h):
        """Bild flaechenfuellend einsetzen: auf das Seitenverhaeltnis des Rahmens zuschneiden."""
        from PIL import Image
        b = Image.open(BytesIO(roh))
        q, z = b.width / b.height, w / h
        p = s.shapes.add_picture(BytesIO(roh), x, y, w, h)
        if q > z:
            weg = (1 - z / q) / 2
            p.crop_left = p.crop_right = weg
        else:
            weg = (1 - q / z) / 2
            p.crop_top = p.crop_bottom = weg
        return p

    def grund(s, hexwert: str):
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = farbe(hexwert)

    def fuss(s, nr: int, dunkel: bool):
        c = "8A97A8" if dunkel else GRAU
        text(s, Inches(0.7), H - Inches(0.55), Inches(8), Inches(0.3), deck, 10, c)
        text(s, B - Inches(1.7), H - Inches(0.55), Inches(1.0), Inches(0.3), f"{nr} / {len(folien)}", 10, c, ausrichtung=PP_ALIGN.RIGHT)

    def kopf(s, t: str, breite=None):
        flaeche(s, Inches(0.7), Inches(0.62), Inches(0.6), Emu(38100), AKZENT)
        text(s, Inches(0.7), Inches(0.8), breite or (B - Inches(1.4)), Inches(1.3), t,
             _schrift(len(t), 32, 24, 60), DUNKEL, KOPF_SCHRIFT, anker=MSO_ANCHOR.TOP, abstand=1.0)

    for i, f in enumerate(folien):
        s = prs.slides.add_slide(leer)
        roh = folienbilder.vorhanden(eid, f["bild"]) if eid and f["bild"] else None
        t = f["titel"] or titel or ""
        typ = f["typ"]
        dunkel = typ in ("titel", "kapitel", "zitat", "abschluss") or (typ == "bild" and roh)
        grund(s, DUNKEL if dunkel else PAPIER)
        if typ == "titel":
            if roh:
                bild(s, roh, Inches(5.6), 0, B - Inches(5.6), H)
            flaeche(s, Inches(0.8), Inches(2.3), Inches(0.9), Emu(57150), AKZENT)
            text(s, Inches(0.8), Inches(2.55), Inches(6.6) if roh else Inches(11.5), Inches(2.6), t, _schrift(len(t), 46, 32, 40),
                 "FFFFFF", KOPF_SCHRIFT, anker=MSO_ANCHOR.TOP, abstand=1.0)
            if f["unter"]:
                text(s, Inches(0.8), Inches(5.1), Inches(6.4) if roh else Inches(11.5), Inches(1.0), f["unter"], 20, "E8C27A")
            text(s, Inches(0.8), H - Inches(0.9), Inches(4), Inches(0.4), time.strftime("%d.%m.%Y"), 12, "8A97A8")
        elif typ == "kapitel":
            text(s, Inches(0.8), Inches(1.6), Inches(4), Inches(2.2), f"{f.get('nr', 1):02d}", 120, AKZENT, KOPF_SCHRIFT)
            text(s, Inches(0.8), Inches(3.9), Inches(11.5), Inches(1.6), t, _schrift(len(t), 44, 30, 40), "FFFFFF", KOPF_SCHRIFT)
            if f["rest"] or f["punkte"]:
                text(s, Inches(0.8), Inches(5.4), Inches(11.5), Inches(1.0), " · ".join((f["rest"] + f["punkte"])[:3]), 18, "B9C3CF")
            if roh:
                bild(s, roh, Inches(8.6), 0, B - Inches(8.6), H)
        elif typ == "aussage":
            flaeche(s, Inches(0.8), Inches(2.2), Emu(57150), Inches(3.0), AKZENT)
            text(s, Inches(1.2), Inches(2.1), Inches(7.4) if roh else Inches(11.0), Inches(2.6), t, _schrift(len(t), 40, 28, 70),
                 DUNKEL, KOPF_SCHRIFT, anker=MSO_ANCHOR.TOP, abstand=1.0)
            beleg = (f["rest"] + f["punkte"])[:2]
            if beleg:
                text(s, Inches(1.2), Inches(4.8), Inches(7.4) if roh else Inches(11.0), Inches(1.6), beleg, 20, GRAU)
            if roh:
                bild(s, roh, Inches(9.0), 0, B - Inches(9.0), H)
        elif typ == "zahlen":
            kopf(s, t)
            zahlen = f["zahlen"][:4] or [("", p) for p in f["punkte"][:4]]
            n = max(1, len(zahlen))
            bw = (B - Inches(1.4) - Inches(0.4) * (n - 1)) / n
            for k, (zahl, was) in enumerate(zahlen):
                x = Inches(0.7) + k * (bw + Inches(0.4))
                flaeche(s, x, Inches(2.6), bw, Inches(3.6), "FFFFFF")
                flaeche(s, x, Inches(2.6), bw, Emu(57150), AKZENT if k % 2 == 0 else AKZENT2)
                text(s, x + Inches(0.35), Inches(3.0), bw - Inches(0.7), Inches(1.4), zahl or "–", _schrift(len(zahl), 54, 34, 6),
                     AKZENT if k % 2 == 0 else AKZENT2, KOPF_SCHRIFT)
                text(s, x + Inches(0.35), Inches(4.45), bw - Inches(0.7), Inches(1.7), was, _schrift(len(was), 17, 13, 70), TEXT)
        elif typ in ("zwei_spalten", "vergleich"):
            kopf(s, t)
            for k, sp in enumerate(f["spalten"][:2]):
                x = Inches(0.7) + k * Inches(6.15)
                rechts = k == 1
                hinter = (AKZENT2 if rechts else "E9E5DC") if typ == "vergleich" else "FFFFFF"
                flaeche(s, x, Inches(2.5), Inches(5.8), Inches(4.0), hinter)
                c = "FFFFFF" if typ == "vergleich" and rechts else DUNKEL
                text(s, x + Inches(0.4), Inches(2.75), Inches(5.0), Inches(0.6), sp["titel"], 20, c if c != DUNKEL else AKZENT2, KOPF_SCHRIFT)
                ps = sp["punkte"][:5]
                text(s, x + Inches(0.4), Inches(3.45), Inches(5.0), Inches(3.0), [[(p, c, False)] for p in ps],
                     _schrift(sum(len(p) for p in ps), 17, 13, 220), abstand=1.1,
                     punkte=AKZENT if not (typ == "vergleich" and rechts) else "E8C27A")
        elif typ == "zitat":
            text(s, Inches(1.0), Inches(0.9), Inches(2), Inches(2), "“", 160, AKZENT, ZITAT_SCHRIFT)
            text(s, Inches(1.6), Inches(2.2), Inches(10.2), Inches(3.2), f["zitat"] or t, _schrift(len(f["zitat"] or t), 34, 22, 140),
                 "FFFFFF", ZITAT_SCHRIFT, kursiv=True, abstand=1.15)
            if f["quelle"]:
                text(s, Inches(1.6), Inches(5.6), Inches(10), Inches(0.6), "— " + f["quelle"], 18, "E8C27A")
        elif typ == "bild" and roh:
            bild(s, roh, 0, 0, B, H)
            flaeche(s, 0, H - Inches(2.3), B, Inches(2.3), DUNKEL, alpha=72)
            text(s, Inches(0.8), H - Inches(2.0), Inches(11.5), Inches(1.4), t, _schrift(len(t), 36, 26, 70), "FFFFFF", KOPF_SCHRIFT,
                 anker=MSO_ANCHOR.MIDDLE)
        elif typ == "abschluss":
            text(s, Inches(0.8), Inches(1.2), Inches(11.5), Inches(1.8), t, _schrift(len(t), 40, 28, 60), "FFFFFF", KOPF_SCHRIFT)
            flaeche(s, Inches(0.8), Inches(3.05), Inches(0.9), Emu(57150), AKZENT)
            ps = (f["punkte"] + f["rest"])[:4]
            for k, p in enumerate(ps):
                y = Inches(3.5) + k * Inches(0.85)
                text(s, Inches(0.8), y, Inches(0.8), Inches(0.7), f"{k + 1}", 26, AKZENT, KOPF_SCHRIFT)
                text(s, Inches(1.5), y + Inches(0.05), Inches(10.5), Inches(0.75), p, _schrift(len(p), 20, 15, 90), "E6EAF0")
        elif typ == "quellen":
            kopf(s, "Quellen")
            qs = f.get("quellen") or []
            haelfte = (len(qs) + 1) // 2
            for k, teil in enumerate((qs[:haelfte], qs[haelfte:])):
                if teil:
                    text(s, Inches(0.7) + k * Inches(6.1), Inches(2.0), Inches(5.8), Inches(4.6),
                         [[(f"{n + 1 + k * haelfte}. {q[0][:80]}", TEXT, False), (f"  {re.sub(r'^https?://(www[.])?', '', q[1]).split('/')[0]}", GRAU, False)]
                          for n, q in enumerate(teil)], 11, abstand=1.05)
        else:  # punkte (und bild ohne gerechnetes Bild)
            breit = Inches(7.0) if roh else B - Inches(1.4)
            kopf(s, t, breit)
            ps = (f["punkte"] or f["rest"])[:6]
            text(s, Inches(0.7), Inches(2.5), breit, Inches(4.2), ps,
                 _schrift(sum(len(p) for p in ps), 21, 15, 260), abstand=1.1, punkte=AKZENT)
            if roh:
                bild(s, roh, Inches(8.3), 0, B - Inches(8.3), H)
        if typ not in ("titel", "bild"):
            fuss(s, i + 1, dunkel)
        if f["notizen"]:
            s.notes_slide.notes_text_frame.text = "\n".join(f["notizen"])
    puffer = BytesIO()
    prs.save(puffer)
    return puffer.getvalue()


# ------------------------------------------------------------ Vorschau (HTML)
def _inline(s: str) -> str:
    from .stab import _md_inline
    return _md_inline(s)


def zu_html(titel: str, md: str, eid: str | None = None) -> str:
    """Dieselben Folien wie in PowerPoint, auf einer 1280x720-Buehne."""
    from . import folienbilder
    folien = parse(md)
    deck = (folien[0]["titel"] if folien and folien[0]["typ"] == "titel" and folien[0]["titel"] else titel or "")[:80]
    teile = []
    for i, f in enumerate(folien):
        roh = folienbilder.vorhanden(eid, f["bild"]) if eid and f["bild"] else None
        bild = (f"<img class='bild' alt='' src='{folienbilder.daten_uri(roh)}'>" if roh
                else "<div class='bild kommt'>Bild wird gerechnet …</div>" if f["bild"] else "")
        t = _inline(f["titel"] or titel or "")
        typ = f["typ"]
        dunkel = typ in ("titel", "kapitel", "zitat", "abschluss") or (typ == "bild" and roh)
        innen = ""
        if typ == "titel":
            innen = f"<div class='strich'></div><h1>{t}</h1>" + (f"<p class='unter'>{_inline(f['unter'])}</p>" if f["unter"] else "") \
                    + f"<p class='datum'>{time.strftime('%d.%m.%Y')}</p>"
        elif typ == "kapitel":
            innen = f"<div class='knr'>{f.get('nr', 1):02d}</div><h1>{t}</h1>" + \
                    (f"<p class='unter hell'>{' · '.join(_inline(x) for x in (f['rest'] + f['punkte'])[:3])}</p>" if f["rest"] or f["punkte"] else "")
        elif typ == "aussage":
            innen = f"<div class='aussage'><h1>{t}</h1>" + "".join(f"<p>{_inline(x)}</p>" for x in (f["rest"] + f["punkte"])[:2]) + "</div>"
        elif typ == "zahlen":
            zahlen = f["zahlen"][:4] or [("", p) for p in f["punkte"][:4]]
            innen = f"<div class='strich'></div><h2>{t}</h2><div class='zahlen'>" + "".join(
                f"<div class='karte k{k % 2}'><b>{htmlmod.escape(z) or '–'}</b><span>{_inline(w)}</span></div>" for k, (z, w) in enumerate(zahlen)) + "</div>"
        elif typ in ("zwei_spalten", "vergleich"):
            innen = f"<div class='strich'></div><h2>{t}</h2><div class='spalten {typ}'>" + "".join(
                f"<div class='sp sp{k}'><h3>{_inline(sp['titel'])}</h3><ul>" + "".join(f"<li>{_inline(p)}</li>" for p in sp["punkte"][:5]) + "</ul></div>"
                for k, sp in enumerate(f["spalten"][:2])) + "</div>"
        elif typ == "zitat":
            innen = f"<div class='qz'>“</div><blockquote>{_inline(f['zitat'] or f['titel'])}</blockquote>" + \
                    (f"<p class='quelle'>— {_inline(f['quelle'])}</p>" if f["quelle"] else "")
        elif typ == "bild" and roh:
            innen = f"<div class='band'><h2>{t}</h2></div>"
        elif typ == "abschluss":
            innen = f"<h1>{t}</h1><div class='strich'></div><ol>" + "".join(f"<li>{_inline(p)}</li>" for p in (f["punkte"] + f["rest"])[:4]) + "</ol>"
        elif typ == "quellen":
            innen = "<div class='strich'></div><h2>Quellen</h2><ol class='ql'>" + "".join(
                f"<li>{htmlmod.escape(q[0][:80])} <span>{htmlmod.escape(re.sub(r'^https?://(www[.])?', '', q[1]).split('/')[0])}</span></li>"
                for q in f.get("quellen") or []) + "</ol>"
        else:
            innen = f"<div class='strich'></div><h2>{t}</h2><ul class='pk'>" + "".join(f"<li>{_inline(p)}</li>" for p in (f["punkte"] or f["rest"])[:6]) + "</ul>"
        mitbild = bild and typ in ("titel", "kapitel", "aussage", "punkte")
        voll = bild and typ == "bild"
        fuss = "" if typ in ("titel", "bild") else f"<div class='fuss'><span>{htmlmod.escape(deck)}</span><span>{i + 1} / {len(folien)}</span></div>"
        notiz = f"<details class='notiz'><summary>Sprechernotizen</summary>{''.join(f'<p>{_inline(n)}</p>' for n in f['notizen'])}</details>" if f["notizen"] else ""
        teile.append(f"<section class='{typ}{' dunkel' if dunkel else ''}{' mitbild' if mitbild else ''}{' voll' if voll else ''}'>"
                     + (bild if bild and (mitbild or voll) else "") + f"<div class='inhalt'>{innen}</div>{fuss}</section>{notiz}")
    css = (
        f"html,body{{margin:0;background:#E5E2DA;color:#{TEXT};overflow-x:hidden}}"
        ".buehne{width:1280px;transform-origin:0 0;padding:16px 0}"
        f"section{{width:1280px;height:720px;box-sizing:border-box;position:relative;overflow:hidden;background:#{PAPIER};margin:0 0 8px;"
        "border-radius:14px;box-shadow:0 6px 22px rgba(12,26,46,.14);font-family:'Segoe UI',system-ui,sans-serif}"
        f"section.dunkel{{background:#{DUNKEL};color:#fff}}"
        ".inhalt{position:absolute;inset:60px 68px 70px 68px;display:flex;flex-direction:column}"
        "section.mitbild .inhalt{right:44%}"
        "section.kapitel.mitbild .inhalt,section.aussage.mitbild .inhalt{right:38%}"
        ".bild{position:absolute;top:0;right:0;bottom:0;width:40%;height:100%;object-fit:cover}"
        "section.titel .bild{width:58%}section.kapitel .bild{width:35%}section.aussage .bild{width:32%}"
        "section.voll .bild{width:100%}"
        ".bild.kommt{display:flex;align-items:center;justify-content:center;background:#E9E5DC;color:#8A94A6;font-size:22px}"
        "section.dunkel .bild.kommt{background:#16263A;color:#8A97A8}"
        f".strich{{width:58px;height:4px;background:#{AKZENT};margin:0 0 14px}}"
        f"h1,h2,h3{{font-family:'Segoe UI Semibold','Segoe UI',system-ui,sans-serif;font-weight:600;margin:0;color:#{DUNKEL};line-height:1.08}}"
        "section.dunkel h1,section.dunkel h2{color:#fff}"
        "h2{font-size:42px;margin-bottom:34px}"
        "section.titel .inhalt{justify-content:center}section.titel h1{font-size:60px}"
        ".unter{font-size:26px;color:#E8C27A;margin:24px 0 0}.unter.hell{color:#B9C3CF}.datum{position:absolute;bottom:-30px;font-size:16px;color:#8A97A8}"
        f".knr{{font-size:150px;font-weight:600;color:#{AKZENT};line-height:1;margin-top:40px}}section.kapitel h1{{font-size:56px;margin-top:20px}}"
        f".aussage{{border-left:6px solid #{AKZENT};padding-left:34px;margin-top:110px}}.aussage h1{{font-size:50px}}.aussage p{{font-size:25px;color:#{GRAU};margin:26px 0 0}}"
        f"ul.pk{{list-style:none;padding:0;margin:0;font-size:27px;line-height:1.35}}ul.pk li{{margin:0 0 18px;padding-left:34px;position:relative}}"
        f"ul.pk li:before{{content:'';position:absolute;left:0;top:13px;width:11px;height:11px;background:#{AKZENT}}}"
        ".zahlen{display:flex;gap:28px;flex:1;align-items:stretch;margin-top:10px}"
        f".karte{{flex:1;background:#fff;border-top:5px solid #{AKZENT};padding:30px 28px;display:flex;flex-direction:column;gap:16px;max-height:340px}}"
        f".karte.k1{{border-top-color:#{AKZENT2}}}.karte b{{font-size:66px;color:#{AKZENT};line-height:1}}.karte.k1 b{{color:#{AKZENT2}}}.karte span{{font-size:21px;line-height:1.35}}"
        ".spalten{display:flex;gap:30px;flex:1}.sp{flex:1;background:#fff;padding:26px 32px;max-height:380px}"
        f".sp h3{{font-size:26px;color:#{AKZENT2};margin-bottom:14px}}.sp ul{{padding-left:22px;margin:0;font-size:22px;line-height:1.35}}.sp li{{margin:0 0 10px}}"
        f".vergleich .sp0{{background:#E9E5DC}}.vergleich .sp1{{background:#{AKZENT2};color:#fff}}.vergleich .sp1 h3{{color:#fff}}"
        f".qz{{font-family:Georgia,serif;font-size:200px;color:#{AKZENT};line-height:.8;height:110px}}"
        "blockquote{font-family:Georgia,serif;font-style:italic;font-size:42px;line-height:1.3;margin:0 0 0 60px}"
        ".quelle{font-size:22px;color:#E8C27A;margin:30px 0 0 60px}"
        f".band{{position:absolute;left:-68px;right:-68px;bottom:-70px;background:rgba(11,24,38,.72);padding:40px 68px}}.band h2{{color:#fff;margin:0}}"
        "section.abschluss h1{font-size:50px;margin:30px 0 26px}section.abschluss ol{font-size:27px;line-height:1.4;padding-left:42px;margin:20px 0 0;color:#E6EAF0}"
        f"section.abschluss li{{margin:0 0 16px}}section.abschluss li::marker{{color:#{AKZENT};font-weight:700}}"
        ".fuss{position:absolute;left:68px;right:68px;bottom:24px;display:flex;justify-content:space-between;font-size:14px;color:#8A97A8}"
        "ol.ql{columns:2;column-gap:48px;font-size:15px;line-height:1.35;margin:0;padding-left:24px}ol.ql li{margin:0 0 8px;break-inside:avoid}ol.ql span{color:#8A94A6}"
        ".notiz{width:1280px;margin:0 0 22px;font:15px/1.5 system-ui,sans-serif;color:#5B6878}.notiz summary{cursor:pointer}.notiz p{margin:4px 0}")
    return ("<!doctype html><meta charset='utf-8'><meta name='viewport' content='width=device-width'>"
            f"<title>{htmlmod.escape(titel)}</title><style>{css}</style><div class='buehne'>" + "".join(teile) + "</div><script>"
            "function passen(){var b=document.querySelector('.buehne');var w=document.documentElement.clientWidth;"
            "var s=Math.min(1.6,(w-32)/1280);b.style.transform='scale('+s+')';b.style.marginLeft=((w-1280*s)/2)+'px';"
            "document.body.style.height=(b.offsetHeight*s)+'px';}"
            # Zu viel Text: die Schrift im Inhaltsbereich schrittweise verkleinern, statt ueberlaufen zu lassen.
            "document.querySelectorAll('section .inhalt').forEach(function(f){var g=100;"
            "while(f.scrollHeight>f.clientHeight+2&&g>60){g-=5;f.style.fontSize=g+'%';f.querySelectorAll('h1,h2,ul,ol,blockquote,p,.karte span,.sp ul').forEach(function(e){e.style.zoom=g/100;});}});"
            "passen();window.addEventListener('resize',passen);</script>")
