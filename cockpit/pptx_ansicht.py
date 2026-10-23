"""
POWERPOINT ALS FOLIEN ANSEHEN (23.10.2026).

Veiko: eine PowerPoint aus dem Explorer "scheint er nicht als wirklich PowerPoint gut
oeffnen zu koennen" -- und die FR5-Magazinausgabe steht im HOCHFORMAT. Bis dahin
zeigte das Fenster nur den gelesenen Text (anhang._pptx_lesen) als Word-Seite:
Woerter ohne Bilder, ohne Anordnung, ohne Format.

Auf dem Server gibt es kein LibreOffice und keine Moeglichkeit, es nachzuinstallieren
(vveadmin darf kein apt). Deshalb wird die Folie hier SELBST gelesen: python-pptx
liefert Groesse und Lage jeder Form, Text mit Schriftgroesse, Fett, Farbe, Bilder mit
Zuschnitt. Daraus entsteht eine Beschreibung (JSON), die das Cockpit als absolut
gesetzte Folie im RICHTIGEN Seitenverhaeltnis zeichnet (datei.js, FolienAnsicht).

Bewusst KEIN HTML aus dem Server: der Text kommt aus einer fremden Datei, und Preact
setzt ihn als Text, nie als Markup. Bilder laufen ueber /dateien/{id}/folienmedium --
verkleinert und zwischengespeichert, sonst waere die 41-MB-Ausgabe ein 41-MB-Abruf.

Was NICHT nachgebaut wird und so benannt ist: Diagramme, SmartArt, Animationen,
WordArt-Effekte, EMF/WMF-Grafiken (der Browser kann sie nicht zeigen). Dort steht ein
grauer Platzhalter mit dem Hinweis -- nichts wird erfunden.
"""
from __future__ import annotations

import hashlib
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any

from .konfig import DATEN

EMU_PT = 12700
CACHE = DATEN / "cache" / "folien"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
BROWSER_BILD = {"png", "jpg", "jpeg", "gif", "bmp", "webp", "svg", "tif", "tiff"}
_SCHEMA_NAMEN = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}
_merk: dict[str, dict[str, Any]] = {}


def _pt(emu: Any) -> float:
    return round((int(emu or 0)) / EMU_PT, 2)


# ------------------------------------------------------------------ Farben
def _thema(master: Any) -> dict[str, str]:
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    try:
        teil = master.part.part_related_by(RT.THEME)
        from lxml import etree
        x = etree.fromstring(teil.blob)
    except Exception:  # noqa: BLE001
        return {}
    farben: dict[str, str] = {}
    schema = x.find(f".//{A}clrScheme")
    if schema is None:
        return farben
    for el in schema:
        name = el.tag.replace(A, "")
        kind = el[0] if len(el) else None
        if kind is None:
            continue
        wert = kind.get("val") if kind.tag == f"{A}srgbClr" else kind.get("lastClr")
        if wert:
            farben[name] = wert.upper()
    return farben


def _mischen(hexwert: str, el: Any) -> tuple[str, float]:
    """Wendet lumMod/lumOff/shade/tint/alpha (das Noetigste) auf eine Farbe an."""
    r, g, b = (int(hexwert[i:i + 2], 16) for i in (0, 2, 4))
    alpha = 1.0
    for m in el:
        name, v = m.tag.replace(A, ""), int(m.get("val") or 0) / 100000
        if name == "lumMod":
            r, g, b = (c * v for c in (r, g, b))
        elif name == "lumOff":
            r, g, b = (c + 255 * v for c in (r, g, b))
        elif name == "shade":
            r, g, b = (c * v for c in (r, g, b))
        elif name == "tint":
            r, g, b = (c + (255 - c) * (1 - v) for c in (r, g, b))
        elif name == "alpha":
            alpha = v
    r, g, b = (max(0, min(255, int(round(c)))) for c in (r, g, b))
    return f"#{r:02X}{g:02X}{b:02X}", alpha


def _farbe(fuellung: Any, thema: dict[str, str]) -> tuple[str, float] | None:
    """Farbe aus einem a:solidFill (oder einem Element, das eine Farbe direkt traegt)."""
    if fuellung is None:
        return None
    for el in fuellung if fuellung.tag == f"{A}solidFill" else [fuellung]:
        tag = el.tag.replace(A, "")
        if tag == "srgbClr" and el.get("val"):
            return _mischen(el.get("val").upper(), el)
        if tag == "schemeClr":
            name = _SCHEMA_NAMEN.get(el.get("val"), el.get("val"))
            if name in thema:
                return _mischen(thema[name], el)
        if tag == "sysClr":
            return _mischen((el.get("lastClr") or "000000").upper(), el)
        if tag == "prstClr":
            return _mischen({"black": "000000", "white": "FFFFFF", "red": "FF0000", "blue": "0000FF",
                             "green": "008000", "gray": "808080", "yellow": "FFFF00"}.get(el.get("val"), "000000"), el)
    return None


def _hex(f: tuple[str, float] | None) -> str | None:
    if not f:
        return None
    hexwert, alpha = f
    if alpha >= 0.999:
        return hexwert
    return f"rgba({int(hexwert[1:3], 16)},{int(hexwert[3:5], 16)},{int(hexwert[5:7], 16)},{round(alpha, 3)})"


# ------------------------------------------------------------------ Text
def _stil_ebene(quellen: list[Any], ebene: int) -> dict[str, Any]:
    """Schriftvorgaben je Ebene aus lstStyle/txStyles (spaetere Quelle gewinnt)."""
    aus: dict[str, Any] = {}
    for q in quellen:
        if q is None:
            continue
        lvl = q.find(f"{A}lvl{ebene + 1}pPr")
        if lvl is None:
            continue
        if lvl.get("algn"):
            aus["algn"] = lvl.get("algn")
        # Aufzaehlungszeichen der Vorlage (Inhaltsplatzhalter tragen sie meist nur dort).
        if lvl.find(f"{A}buChar") is not None:
            aus["aufz"] = lvl.find(f"{A}buChar").get("char") or "•"
        elif lvl.find(f"{A}buAutoNum") is not None:
            aus["aufz"] = "1."
        elif lvl.find(f"{A}buNone") is not None:
            aus["aufz"] = None
        if lvl.get("marL"):
            aus["einzug"] = _pt(lvl.get("marL"))
        if lvl.get("indent"):
            aus["erste"] = _pt(lvl.get("indent"))
        d = lvl.find(f"{A}defRPr")
        if d is not None:
            aus.setdefault("rpr", []).append(d)
    return aus


def _lauf_stil(rprs: list[Any], thema: dict[str, str], aus: dict[str, Any]) -> dict[str, Any]:
    for rpr in rprs:
        if rpr is None:
            continue
        if rpr.get("sz"):
            aus["gr"] = int(rpr.get("sz")) / 100
        if rpr.get("b") is not None:
            aus["f"] = rpr.get("b") in ("1", "true")
        if rpr.get("i") is not None:
            aus["k"] = rpr.get("i") in ("1", "true")
        if rpr.get("u") is not None:
            aus["u"] = rpr.get("u") not in ("none", "0")
        if rpr.get("cap") == "all":
            aus["gross"] = True
        if rpr.get("spc"):
            aus["spc"] = int(rpr.get("spc")) / 100
        fl = rpr.find(f"{A}solidFill")
        if fl is not None:
            aus["farbe"] = _hex(_farbe(fl, thema))
        lat = rpr.find(f"{A}latin")
        if lat is not None and lat.get("typeface") and not lat.get("typeface").startswith("+"):
            aus["schrift"] = lat.get("typeface")
    return aus


def _text(txbody: Any, thema: dict[str, str], vorgaben: list[Any]) -> dict[str, Any] | None:
    if txbody is None:
        return None
    bpr = txbody.find(f"{A}bodyPr")
    lst = txbody.find(f"{A}lstStyle")
    absaetze = []
    for p in txbody.findall(f"{A}p"):
        ppr = p.find(f"{A}pPr")
        ebene = int(ppr.get("lvl") or 0) if ppr is not None else 0
        st = _stil_ebene(vorgaben + [lst], ebene)
        grund = _lauf_stil(st.get("rpr", []), thema, {})
        ausr = (ppr.get("algn") if ppr is not None and ppr.get("algn") else st.get("algn")) or "l"
        abs_: dict[str, Any] = {"ausr": {"ctr": "c", "r": "r", "just": "j", "dist": "j"}.get(ausr, "l"), "laeufe": []}
        for schl in ("aufz", "einzug", "erste"):
            if st.get(schl) is not None:
                abs_[schl] = st[schl]
        if ppr is not None:
            lz = ppr.find(f"{A}lnSpc/{A}spcPct")
            if lz is not None:
                abs_["zh"] = int(lz.get("val")) / 100000
            lp = ppr.find(f"{A}lnSpc/{A}spcPts")
            if lp is not None:
                abs_["zh_pt"] = int(lp.get("val")) / 100
            for name, schl in (("spcBef", "vor"), ("spcAft", "nach")):
                s = ppr.find(f"{A}{name}/{A}spcPts")
                if s is not None:
                    abs_[schl] = int(s.get("val")) / 100
            if ppr.find(f"{A}buChar") is not None:
                abs_["aufz"] = ppr.find(f"{A}buChar").get("char") or "•"
            elif ppr.find(f"{A}buAutoNum") is not None:
                abs_["aufz"] = "1."
            elif ppr.find(f"{A}buNone") is not None:
                abs_.pop("aufz", None)
            if ppr.get("marL"):
                abs_["einzug"] = _pt(ppr.get("marL"))
            if ppr.get("indent"):
                abs_["erste"] = _pt(ppr.get("indent"))
        for el in p:
            tag = el.tag.replace(A, "")
            if tag == "r":
                t = (el.findtext(f"{A}t") or "")
                stil = _lauf_stil([el.find(f"{A}rPr")], thema, dict(grund))
                abs_["laeufe"].append({"t": t, **stil})
            elif tag == "br":
                abs_["laeufe"].append({"t": "\n", **grund})
            elif tag == "fld":
                stil = _lauf_stil([el.find(f"{A}rPr")], thema, dict(grund))
                abs_["laeufe"].append({"t": el.findtext(f"{A}t") or "", **stil})
        if not abs_["laeufe"]:
            # Leerer Absatz: haelt trotzdem eine Zeile in der Groesse seines Endzeichens.
            stil = _lauf_stil([p.find(f"{A}endParaRPr")], thema, dict(grund))
            abs_["leer"] = stil.get("gr") or grund.get("gr")
        absaetze.append(abs_)
    if not any(a["laeufe"] and "".join(x["t"] for x in a["laeufe"]).strip() for a in absaetze):
        return None
    aus: dict[str, Any] = {"absaetze": absaetze, "anker": "t", "innen": [7.2, 3.6, 7.2, 3.6], "umbruch": True}
    if bpr is not None:
        aus["anker"] = {"ctr": "m", "b": "b"}.get(bpr.get("anchor") or "t", "t")
        aus["innen"] = [_pt(bpr.get(k)) if bpr.get(k) is not None else d
                        for k, d in (("lIns", 7.2), ("tIns", 3.6), ("rIns", 7.2), ("bIns", 3.6))]
        aus["umbruch"] = bpr.get("wrap") != "none"
        if bpr.get("vert") in ("vert", "vert270", "eaVert"):
            aus["senkrecht"] = bpr.get("vert")
        na = bpr.find(f"{A}normAutofit")
        if na is not None and na.get("fontScale"):
            aus["skala"] = int(na.get("fontScale")) / 100000
    return aus


# ------------------------------------------------------------------ Formen
class _Lesen:
    def __init__(self, praes: Any, zipnamen: set[str]):
        self.praes = praes
        self.zipnamen = zipnamen
        self.hinweise: set[str] = set()

    def medium(self, teil: Any, rid: str) -> str | None:
        try:
            ziel = teil.related_part(rid)
        except (KeyError, AttributeError):
            return None
        name = str(ziel.partname).lstrip("/")
        return name if name in self.zipnamen else None

    def master_stile(self, master: Any) -> dict[str, Any]:
        tx = master._element.find(f"{P}txStyles")
        if tx is None:
            return {}
        return {"titel": tx.find(f"{P}titleStyle"), "text": tx.find(f"{P}bodyStyle"), "sonst": tx.find(f"{P}otherStyle")}

    def formen(self, shapes: Any, teil: Any, thema: dict[str, str], stile: dict[str, Any],
               nur_deko: bool = False, umrechnen: Any = None) -> list[dict[str, Any]]:
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        aus: list[dict[str, Any]] = []
        for sh in shapes:
            try:
                if nur_deko and sh.is_placeholder:
                    continue
                x, y, b, h = (sh.left or 0), (sh.top or 0), (sh.width or 0), (sh.height or 0)
                if umrechnen:
                    x, y, b, h = umrechnen(x, y, b, h)
                f: dict[str, Any] = {"x": _pt(x), "y": _pt(y), "b": _pt(b), "h": _pt(h)}
                el = sh._element
                xfrm = el.find(f".//{A}xfrm")
                if xfrm is not None:
                    if xfrm.get("rot"):
                        f["dreh"] = int(xfrm.get("rot")) / 60000
                    if xfrm.get("flipH") == "1":
                        f["spiegelH"] = True
                    if xfrm.get("flipV") == "1":
                        f["spiegelV"] = True
                st = sh.shape_type
                if st == MSO_SHAPE_TYPE.GROUP:
                    gx = el.find(f"{P}grpSpPr/{A}xfrm")
                    if gx is None:
                        continue
                    off, ext = gx.find(f"{A}off"), gx.find(f"{A}ext")
                    choff, chext = gx.find(f"{A}chOff"), gx.find(f"{A}chExt")
                    ox, oy = int(off.get("x")), int(off.get("y"))
                    ex, ey = int(ext.get("cx")) or 1, int(ext.get("cy")) or 1
                    cx0, cy0 = int(choff.get("x")), int(choff.get("y"))
                    cex, cey = int(chext.get("cx")) or 1, int(chext.get("cy")) or 1

                    def inner(ix, iy, ib, ih, _o=(ox, oy, ex, ey, cx0, cy0, cex, cey), _u=umrechnen):
                        ox_, oy_, ex_, ey_, c0x, c0y, cx_, cy_ = _o
                        sx, sy = ex_ / cx_, ey_ / cy_
                        r = (ox_ + (ix - c0x) * sx, oy_ + (iy - c0y) * sy, ib * sx, ih * sy)
                        return _u(*r) if _u else r
                    aus.extend(self.formen(sh.shapes, teil, thema, stile, False, inner))
                    continue
                if st == MSO_SHAPE_TYPE.PICTURE or el.find(f".//{A}blip") is not None and not sh.has_text_frame:
                    blip = el.find(f".//{A}blip")
                    rid = blip.get(f"{R}embed") if blip is not None else None
                    name = self.medium(teil, rid) if rid else None
                    f["art"] = "bild"
                    if name and name.lower().rsplit(".", 1)[-1] in BROWSER_BILD:
                        f["medium"] = name
                    else:
                        f["fehlt"] = "Grafik in einem Format, das der Browser nicht zeigen kann"
                        self.hinweise.add("Manche Grafiken (EMF/WMF) lassen sich im Browser nicht anzeigen.")
                    src = el.find(f".//{A}srcRect")
                    if src is not None:
                        f["schnitt"] = [int(src.get(k) or 0) / 100000 for k in ("l", "t", "r", "b")]
                    aus.append(f)
                    continue
                if st == MSO_SHAPE_TYPE.TABLE or el.find(f".//{A}tbl") is not None:
                    tbl = el.find(f".//{A}tbl")
                    spalten = [_pt(g.get("w")) for g in tbl.findall(f"{A}tblGrid/{A}gridCol")]
                    zeilen = []
                    for tr in tbl.findall(f"{A}tr"):
                        zellen = []
                        for tc in tr.findall(f"{A}tc"):
                            tcpr = tc.find(f"{A}tcPr")
                            z = {"text": _text(tc.find(f"{A}txBody"), thema, [stile.get("sonst")]),
                                 "grund": _hex(_farbe(tcpr.find(f"{A}solidFill"), thema)) if tcpr is not None else None}
                            if tc.get("gridSpan"):
                                z["spanne"] = int(tc.get("gridSpan"))
                            if tc.get("hMerge") or tc.get("vMerge"):
                                z["verbunden"] = True
                            zellen.append(z)
                        zeilen.append({"h": _pt(tr.get("h")), "zellen": zellen})
                    f.update({"art": "tabelle", "spalten": spalten, "zeilen": zeilen})
                    aus.append(f)
                    continue
                if el.tag == f"{P}graphicFrame":
                    f.update({"art": "platzhalter", "fehlt": "Diagramm oder SmartArt -- im Cockpit nicht nachgebaut"})
                    self.hinweise.add("Diagramme und SmartArt erscheinen als grauer Platzhalter.")
                    aus.append(f)
                    continue
                # Form, Textfeld, Linie, Platzhalter
                f["art"] = "linie" if el.tag == f"{P}cxnSp" else "form"
                sppr = el.find(f"{P}spPr")
                if sppr is not None:
                    fl = sppr.find(f"{A}solidFill")
                    if fl is not None:
                        f["fuellung"] = _hex(_farbe(fl, thema))
                    gr = sppr.find(f"{A}gradFill")
                    if gr is not None:
                        stops = [_hex(_farbe(gs[0] if len(gs) else None, thema)) for gs in gr.findall(f"{A}gsLst/{A}gs")]
                        stops = [s for s in stops if s]
                        if stops:
                            f["verlauf"] = stops
                    bf = sppr.find(f"{A}blipFill/{A}blip")
                    if bf is not None and bf.get(f"{R}embed"):
                        name = self.medium(teil, bf.get(f"{R}embed"))
                        if name:
                            f["bildfuellung"] = name
                    ln = sppr.find(f"{A}ln")
                    if ln is not None and ln.find(f"{A}noFill") is None:
                        farbe = _hex(_farbe(ln.find(f"{A}solidFill"), thema))
                        if farbe:
                            f["rand"] = {"farbe": farbe, "breite": max(0.5, _pt(ln.get("w") or 9525))}
                    geo = sppr.find(f"{A}prstGeom")
                    if geo is not None:
                        if geo.get("prst") in ("ellipse", "circle"):
                            f["rund"] = "50%"
                        elif geo.get("prst") in ("roundRect", "flowChartAlternateProcess"):
                            f["rund"] = "12%"
                # Linien ohne eigene Linienfarbe nehmen den Stil der Form (style/lnRef).
                if f["art"] == "linie" and "rand" not in f:
                    ref = el.find(f"{P}style/{A}lnRef")
                    farbe = _hex(_farbe(ref[0] if ref is not None and len(ref) else None, thema))
                    f["rand"] = {"farbe": farbe or "#000000", "breite": 1}
                if "fuellung" not in f and "verlauf" not in f and el.find(f"{P}spPr/{A}noFill") is None and not sh.has_text_frame:
                    ref = el.find(f"{P}style/{A}fillRef")
                    if ref is not None and len(ref) and ref.get("idx") not in ("0", None):
                        f["fuellung"] = _hex(_farbe(ref[0], thema))
                if sh.has_text_frame:
                    art = "sonst"
                    if sh.is_placeholder:
                        typ = str(sh.placeholder_format.type or "")
                        art = "titel" if ("TITLE" in typ and "SUB" not in typ) or "CENTER_TITLE" in typ else "text"
                    vorg = [stile.get(art)]
                    # Platzhalter erben Schrift aus dem Layout-Platzhalter.
                    if sh.is_placeholder:
                        try:
                            lay = sh._base_placeholder
                            if lay is not None:
                                lst = lay._element.find(f"{P}txBody/{A}lstStyle")
                                vorg.append(lst)
                        except Exception:  # noqa: BLE001
                            pass
                    t = _text(el.find(f"{P}txBody"), thema, vorg)
                    if t:
                        if art == "titel" and not any(l.get("gr") for a in t["absaetze"] for l in a["laeufe"]):
                            t["grund_gr"] = 32
                        f["text"] = t
                if f["art"] == "form" and not any(k in f for k in ("fuellung", "verlauf", "bildfuellung", "rand", "text")):
                    continue
                aus.append(f)
            except Exception as e:  # noqa: BLE001 -- eine fremde Form darf die Folie nicht kippen
                self.hinweise.add(f"Eine Form ließ sich nicht lesen ({type(e).__name__}).")
        return aus

    def hintergrund(self, teile: list[Any], thema: dict[str, str]) -> dict[str, Any]:
        for t in teile:
            bg = t._element.find(f"{P}cSld/{P}bg")
            if bg is None:
                continue
            bgpr = bg.find(f"{P}bgPr")
            if bgpr is not None:
                fl = bgpr.find(f"{A}solidFill")
                if fl is not None:
                    return {"farbe": _hex(_farbe(fl, thema))}
                bl = bgpr.find(f"{A}blipFill/{A}blip")
                if bl is not None:
                    name = self.medium(t.part, bl.get(f"{R}embed"))
                    if name:
                        return {"bild": name}
                gr = bgpr.find(f"{A}gradFill")
                if gr is not None:
                    stops = [_hex(_farbe(gs[0] if len(gs) else None, thema)) for gs in gr.findall(f"{A}gsLst/{A}gs")]
                    if [s for s in stops if s]:
                        return {"verlauf": [s for s in stops if s]}
            ref = bg.find(f"{P}bgRef")
            if ref is not None and len(ref):
                return {"farbe": _hex(_farbe(ref[0], thema))}
        return {"farbe": "#FFFFFF"}


def folien(pfad: str) -> dict[str, Any]:
    """Beschreibung aller Folien (zwischengespeichert je Datei und Aenderungszeit)."""
    from pptx import Presentation
    p = Path(pfad)
    schl = f"{pfad}:{p.stat().st_mtime}"
    if schl in _merk:
        return _merk[schl]
    praes = Presentation(str(p))
    with zipfile.ZipFile(p) as z:
        namen = set(z.namelist())
    les = _Lesen(praes, namen)
    liste = []
    for i, s in enumerate(praes.slides):
        layout = s.slide_layout
        master = layout.slide_master
        thema = _thema(master)
        stile = les.master_stile(master)
        formen: list[dict[str, Any]] = []
        # Deko aus Master und Layout liegt UNTER der Folie -- es sei denn, die Folie blendet sie aus.
        if s._element.get("showMasterSp") != "0":
            if layout._element.get("showMasterSp") != "0":
                formen += les.formen(master.shapes, master.part, thema, stile, nur_deko=True)
            formen += les.formen(layout.shapes, layout.part, thema, stile, nur_deko=True)
        formen += les.formen(s.shapes, s.part, thema, stile)
        notiz = ""
        if s.has_notes_slide:
            notiz = (s.notes_slide.notes_text_frame.text if s.notes_slide.notes_text_frame else "").strip()
        if s._element.get("show") == "0":
            ausgeblendet = True
        else:
            ausgeblendet = False
        liste.append({"nr": i + 1, "hintergrund": les.hintergrund([s, layout, master], thema), "formen": formen,
                      "notiz": notiz[:4000], "ausgeblendet": ausgeblendet})
    aus = {"breite": _pt(praes.slide_width), "hoehe": _pt(praes.slide_height), "folien": liste,
           "hinweise": sorted(les.hinweise)}
    _merk.clear()   # nur die zuletzt angesehene Datei im Speicher halten
    _merk[schl] = aus
    return aus


def medium(pfad: str, name: str, breite: int = 1600) -> tuple[bytes, str]:
    """Ein Bild aus der pptx, verkleinert und auf der Platte zwischengespeichert.
    Transparenz bleibt (PNG), sonst JPEG -- Logos auf farbigem Grund brauchen ihren Alphakanal."""
    from PIL import Image
    p = Path(pfad)
    breite = max(200, min(2400, int(breite or 1600)))
    schl = hashlib.sha1(f"{pfad}:{p.stat().st_mtime}:{name}:{breite}".encode()).hexdigest()
    CACHE.mkdir(parents=True, exist_ok=True)
    for endung, mime in (("png", "image/png"), ("jpg", "image/jpeg"), ("svg", "image/svg+xml"), ("gif", "image/gif")):
        c = CACHE / f"{schl}.{endung}"
        if c.is_file():
            return c.read_bytes(), mime
    with zipfile.ZipFile(p) as z:
        if not name.startswith("ppt/") or name not in z.namelist():
            raise KeyError(name)
        roh = z.read(name)
    endung = name.lower().rsplit(".", 1)[-1]
    if endung in ("svg", "gif"):
        # SVG wird als <img> gezeigt -- dort laufen keine Skripte; GIF bleibt bewegt.
        (CACHE / f"{schl}.{endung}").write_bytes(roh)
        return roh, "image/svg+xml" if endung == "svg" else "image/gif"
    with Image.open(BytesIO(roh)) as b:
        b.load()
        if max(b.size) > breite:
            b.thumbnail((breite, breite))
        aus = BytesIO()
        if b.mode in ("RGBA", "LA", "P") and (b.mode != "P" or "transparency" in b.info):
            b.convert("RGBA").save(aus, "PNG", optimize=True)
            endung, mime = "png", "image/png"
        else:
            b.convert("RGB").save(aus, "JPEG", quality=85, optimize=True)
            endung, mime = "jpg", "image/jpeg"
    daten = aus.getvalue()
    (CACHE / f"{schl}.{endung}").write_bytes(daten)
    return daten, mime
